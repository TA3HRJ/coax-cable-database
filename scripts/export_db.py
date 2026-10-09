# -*- coding: utf-8 -*-
"""
TA3HX_Coax_Database.xlsx -> export/csv/*.csv + export/coax.sqlite
                          + site/data/cables.min.json (web sitesi verisi)
                          + tests/fixtures/excel_calculator.json (eşlik testi için Excel hesaplayıcı sonuçları)
                          + kaynaklar/README.md (kaynak dizini)

Excel'de hesaplanmış (kaydedilmiş) çalışma kitabının değerlerini okur; formül sonuçları
Excel'in son hesapladığı değerlerdir. Excel'de değişiklik yaptıysanız önce kaydedin.

Kullanım:  python -X utf8 scripts/export_db.py [xlsx_yolu]
"""
import csv
import datetime
import json
import os
import sqlite3
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sanitize_xlsx import sanitize  # noqa: E402
from cables_extra import SIZE_CLASSES, FEATURED_COUNT  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "TA3HX_Coax_Database.xlsx")
OUT = os.path.join(ROOT, "export")
CSV_DIR = os.path.join(OUT, "csv")
DB = os.path.join(OUT, "coax.sqlite")
WEB_JSON = os.path.join(ROOT, "site", "data", "cables.min.json")
FIXTURE = os.path.join(ROOT, "tests", "fixtures", "excel_calculator.json")
SOURCES_MD = os.path.join(ROOT, "kaynaklar", "README.md")

# cables sütunu -> JSON kısa anahtarı (boş değerler yazılmaz)
WEB_CABLE_KEYS = [
    ("cable_id", "id"), ("display_name", "n"), ("short_name", "s"), ("manufacturer", "m"), ("part_number", "pn"),
    ("family", "fam"), ("size_class", "cls"), ("popular_rank", "pop"), ("standard_type", "std"), ("value_type", "val"),
    ("impedance_ohm", "z"), ("impedance_tol_ohm", "ztol"), ("od_mm", "od"), ("inner_material", "icm"),
    ("inner_construction", "icc"), ("inner_od_mm", "icd"), ("dielectric_material", "dm"), ("dielectric_od_mm", "dd"),
    ("shield", "sh"), ("braid_coverage_pct", "bc"), ("jacket", "jk"), ("velocity_factor", "vf"), ("capacitance_pf_m", "cap"),
    ("inductance_uh_m", "ind"), ("dcr_inner_ohm_km", "dci"), ("dcr_outer_ohm_km", "dco"), ("shielding_db", "se"),
    ("fmax_mhz", "fmax"), ("peak_power_kw", "pk"), ("voltage_v", "v"), ("voltage_note", "vn"),
    ("bend_radius_single_mm", "br1"), ("bend_radius_repeated_mm", "br2"), ("weight_kg_km", "w"), ("tensile_kg", "ten"),
    ("temp_min_c", "tmin"), ("temp_max_c", "tmax"), ("outdoor_uv", "out"), ("direct_burial", "bur"), ("fire_class", "fire"),
    ("connectors", "con"), ("k1_pub_db_100ft", "k1p"), ("k2_pub_db_100ft", "k2p"), ("primary_source_id", "src"),
    ("notes", "note"), ("f_min_data", "fa"), ("f_max_data", "fb"), ("k0_fit", "k0"), ("k1_fit", "k1"), ("k2_fit", "k2"),
    ("fit_rms_pct", "rms"),
]

# Excel tablosu -> (SQL tablo adı, birincil anahtar, yalnızca Excel'e özgü yardımcı sütunlar)
TABLES = [
    ("tblSources", "sources", "source_id", []),
    ("tblCables", "cables", "cable_id", ["att_first_row", "pwr_first_row"]),
    ("tblAttenuation", "attenuation", "point_id", ["ln_f", "ln_att", "w0", "w1", "w2", "one", "dev_sq", "order_ok"]),
    ("tblPower", "power", "power_id", ["ln_f", "ln_p", "order_ok"]),
    ("tblLegacyIndex", "legacy_index", "legacy_key", []),
    ("tblLegacyPoints", "legacy_points", None, []),
    ("tblDataDictionary", "data_dictionary", None, []),
]
FOREIGN_KEYS = {
    "cables": [("primary_source_id", "sources", "source_id")],
    "attenuation": [("cable_id", "cables", "cable_id"), ("source_id", "sources", "source_id")],
    "power": [("cable_id", "cables", "cable_id"), ("source_id", "sources", "source_id")],
    "legacy_index": [("cable_id", "cables", "cable_id")],
    "legacy_points": [("legacy_key", "legacy_index", "legacy_key"), ("cable_id", "cables", "cable_id")],
    "ham_bands": [("cable_id", "cables", "cable_id")],
}
INDEXES = {"attenuation": ["cable_id"], "power": ["cable_id"], "legacy_points": ["legacy_key"], "ham_bands": ["cable_id"]}


def read_table(wb, name):
    for ws in wb.worksheets:
        if name in ws.tables:
            rows = list(ws[ws.tables[name].ref])
            hdr = [c.value for c in rows[0]]
            data = [[c.value for c in r] for r in rows[1:]]
            return hdr, data
    raise KeyError(name)


def clean(v):
    if isinstance(v, str):
        if v.startswith("#"):
            raise ValueError(f"Excel hata değeri: {v}")
        return v if v != "" else None
    return v


def ham_bands(wb):
    ws = wb["Ham_Bands"]
    hdr = ["cable_id", "band", "freq_mhz", "att_db_100m", "att_db_100ft", "extrapolated"]
    out = []
    ncol_band = 12
    for r in range(6, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if not cid:
            continue
        fmin, fmax = ws.cell(r, 6).value, ws.cell(r, 7).value
        for k in range(ncol_band):
            v = ws.cell(r, 10 + k).value
            if not isinstance(v, (int, float)):
                continue
            f = ws.cell(5, 10 + k).value
            out.append([cid, ws.cell(4, 10 + k).value, f, v, v * 0.3048, int(f < fmin or f > fmax)])
    return hdr, out


def sql_type(values):
    vals = [v for v in values if v is not None]
    if vals and all(isinstance(v, bool) for v in vals):
        return "INTEGER"
    if vals and all(isinstance(v, int) and not isinstance(v, bool) for v in vals):
        return "INTEGER"
    if vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
        return "REAL"
    return "TEXT"


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return format(v, ".10g")
    return str(v)


def main():
    sanitize(SRC)   # yerel klasör yolu ve kişi adı gibi üst verileri temizle (hücrelere dokunmaz)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    os.makedirs(CSV_DIR, exist_ok=True)
    if os.path.exists(DB):
        os.remove(DB)
    con = sqlite3.connect(DB)
    con.execute("PRAGMA foreign_keys = ON")
    datasets = []
    for xl, sq, pk, drop in TABLES:
        hdr, data = read_table(wb, xl)
        keep = [i for i, h in enumerate(hdr) if h not in drop]
        hdr = [hdr[i] for i in keep]
        data = [[clean(r[i]) for i in keep] for r in data]
        datasets.append((sq, pk, hdr, data))
    hdr, data = ham_bands(wb)
    datasets.append(("ham_bands", None, hdr, data))
    # veri sözlüğü: tablo adlarını SQL adlarına çevir, çıkarılan yardımcı sütunları sil, ham_bands'i ekle
    dd = next(d for d in datasets if d[0] == "data_dictionary")
    xl_to_sql = {xl[3:].lower(): sq for xl, sq, _, _ in TABLES}   # "tblCables" -> "cables"
    dropped = {(sq, c) for _, sq, _, drop in TABLES for c in drop}
    rows = []
    for t, col, desc, unit, kind in dd[3]:
        t = xl_to_sql.get(t.lower(), t.lower())
        if (t, col) not in dropped:
            rows.append([t, col, desc, unit, kind])
    dd[3][:] = rows
    for col, desc, unit in (("cable_id", "cables.cable_id", ""), ("band", "Amatör bandı", ""), ("freq_mhz", "Bant temsil frekansı", "MHz"),
                            ("att_db_100m", "Zayıflama (log-log interpolasyon)", "dB/100m"), ("att_db_100ft", "Zayıflama", "dB/100ft"),
                            ("extrapolated", "1 = datasheet frekans aralığı dışında", "")):
        dd[3].append(["ham_bands", col, desc, unit, "formül"])

    order = ["sources", "cables", "attenuation", "power", "legacy_index", "legacy_points", "data_dictionary", "ham_bands"]
    datasets.sort(key=lambda d: order.index(d[0]))
    for sq, pk, hdr, data in datasets:
        cols = []
        for i, h in enumerate(hdr):
            t = sql_type([r[i] for r in data])
            cols.append(f'"{h}" {t}' + (" PRIMARY KEY" if h == pk else ""))
        for c, rt, rc in FOREIGN_KEYS.get(sq, []):
            cols.append(f'FOREIGN KEY("{c}") REFERENCES {rt}("{rc}")')
        con.execute(f'CREATE TABLE {sq} ({", ".join(cols)})')
        con.executemany(f'INSERT INTO {sq} VALUES ({",".join("?" * len(hdr))})',
                        [[int(v) if isinstance(v, bool) else v for v in r] for r in data])
        for c in INDEXES.get(sq, []):
            con.execute(f'CREATE INDEX ix_{sq}_{c} ON {sq}("{c}")')
        with open(os.path.join(CSV_DIR, f"{sq}.csv"), "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(hdr)
            for r in data:
                w.writerow([fmt(v) for v in r])
        print(f"{sq:16s} {len(data):5d} satır, {len(hdr):2d} sütun")
    con.commit()
    bad = con.execute("PRAGMA foreign_key_check").fetchall()
    assert not bad, bad
    con.execute("VACUUM")
    con.close()
    print("yazıldı:", DB)
    tables = {sq: (hdr, data) for sq, _, hdr, data in datasets}
    write_web_json(tables, wb)
    write_calculator_fixture(wb, tables)
    write_sources_index(tables)


def sig(x, n=6):
    return float(f"{x:.{n}g}") if isinstance(x, float) else x


def rows_as_dicts(tables, name):
    hdr, data = tables[name]
    return [dict(zip(hdr, r)) for r in data]


def write_web_json(tables, wb):
    """Sitenin yüklediği tek dosya: kablolar + sıralı ölçüm noktaları + kaynaklar."""
    att, pwr = {}, {}
    for r in rows_as_dicts(tables, "attenuation"):
        att.setdefault(r["cable_id"], []).append([r["freq_mhz"], sig(r["att_db_100m"]), r["source_id"]])
    for r in rows_as_dicts(tables, "power"):
        pwr.setdefault(r["cable_id"], []).append([r["freq_mhz"], sig(r["power_w"]), r["source_id"]])
    cables = []
    for r in rows_as_dicts(tables, "cables"):
        c = {}
        for col, key in WEB_CABLE_KEYS:
            v = r.get(col)
            if v is None or v == "":
                continue
            c[key] = (v == "Evet") if col == "standard_type" else sig(v)
        c["a"] = sorted(att[r["cable_id"]])
        if r["cable_id"] in pwr:
            c["p"] = sorted(pwr[r["cable_id"]])
        cables.append(c)
    srcs = {}
    for r in rows_as_dicts(tables, "sources"):
        srcs[r["source_id"]] = {k: r[col] for col, k in (("publisher", "p"), ("document_title", "t"), ("revision", "r"),
                                                          ("document_date", "d"), ("host", "h"), ("url", "u"),
                                                          ("notes", "x")) if r[col] not in (None, "", "-")}
    hb = wb["Ham_Bands"]
    bands = [[hb.cell(4, 10 + k).value, hb.cell(5, 10 + k).value] for k in range(12)]
    doc = {
        "v": datetime.date.today().isoformat(),
        "bands": bands,
        "classes": [{"k": k, "tr": tr, "en": en, "use_tr": ut, "use_en": ue} for k, _, tr, en, ut, ue in SIZE_CLASSES],
        "featured": FEATURED_COUNT,
        "src": srcs,
        "cables": cables,
    }
    os.makedirs(os.path.dirname(WEB_JSON), exist_ok=True)
    with open(WEB_JSON, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, separators=(",", ":"))
    print(f"yazıldı: {WEB_JSON} ({os.path.getsize(WEB_JSON) // 1024} KB, {len(cables)} kablo)")


def write_calculator_fixture(wb, tables):
    """Excel Calculator sayfasının girdi/çıktıları - JS hesabının Excel ile aynı olduğunu test etmek için."""
    cs = wb["Calculator"]
    name_to_id = {r["display_name"]: r["cable_id"] for r in rows_as_dicts(tables, "cables")}

    def num(v):
        return v if isinstance(v, (int, float)) else None

    fx = {
        "source": "TA3HX_Coax_Database.xlsx / Calculator (Excel ile hesaplanmış)",
        "inputs": {"cable_id": name_to_id[cs["C5"].value], "f_mhz": cs["C6"].value, "length_m": cs["C7"].value,
                   "p_in_w": cs["C8"].value, "swr": cs["C9"].value},
        "outputs": {"att_db_100m": cs["C13"].value, "matched_loss_db": cs["C15"].value, "total_loss_db": cs["C17"].value,
                    "p_out_w": cs["C18"].value, "swr_tx": cs["C20"].value, "p_max_w": num(cs["C21"].value),
                    "wavelength_m": num(cs["C23"].value)},
        "compare": [],
    }
    for r in range(31, 31 + 100):
        name = cs.cell(r, 2).value
        if not name:
            break
        fx["compare"].append({"cable_id": name_to_id[name], "att_db_100m": cs.cell(r, 3).value,
                              "total_loss_db": cs.cell(r, 4).value, "p_out_w": cs.cell(r, 5).value,
                              "p_max_w": num(cs.cell(r, 7).value)})
    os.makedirs(os.path.dirname(FIXTURE), exist_ok=True)
    with open(FIXTURE, "w", encoding="utf-8") as f:
        json.dump(fx, f, ensure_ascii=False, indent=1)
    print(f"yazıldı: {FIXTURE} ({len(fx['compare'])} kablo)")


def write_sources_index(tables):
    L = ["# Kaynak belgeler", "",
         "Veritabanındaki her değer bu belgelerden birine bağlıdır (`Sources` sayfası / `sources` tablosu). "
         "Bu dosya `scripts/export_db.py` ile üretilir.", "",
         "Üretici PDF'leri üçüncü tarafların telifli belgeleri olduğu için **bu repoda yayımlanmaz** (`.gitignore`). "
         "Aşağıdaki adreslerden indirilip bu klasöre `Dosya` sütunundaki adla kaydedilebilir; `scripts/gen_points.py` "
         "noktaları yeniden çıkarmak için bunlara ihtiyaç duyar. `.txt` dosyaları web sayfalarından alınan verilerin kendi "
         "özetlerimizdir ve repoda bulunur.", "",
         "Bazı bağlantılar zamanla değişebilir; belge tarihi ve revizyon bilgisi bu yüzden kayıtlıdır.", "",
         "| ID | Yayıncı | Belge | Rev./tarih | Barındıran | Dosya |", "|---|---|---|---|---|---|"]
    for r in rows_as_dicts(tables, "sources"):
        rd = " / ".join(x for x in (r["revision"], r["document_date"]) if x and x != "-") or "-"
        t = f"[{r['document_title']}]({r['url'].replace(' ', '%20')})" if r["url"] else r["document_title"]
        f = os.path.basename(r["local_file"]) if r["local_file"] else ""
        L.append(f"| {r['source_id']} | {r['publisher']} | {t} | {rd} | {r['host']} | `{f}` |")
    with open(SOURCES_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print("yazıldı:", SOURCES_MD)


if __name__ == "__main__":
    main()
