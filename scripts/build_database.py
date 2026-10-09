# -*- coding: utf-8 -*-
"""
TA3HX Koaksiyel Kablo Veritabanı - çalışma kitabı üreticisi.

Veri: bu dosyadaki pilot kablolar + scripts/cables_extra.py (kablo bilgileri, elle noktalar)
      + scripts/points_extra.py (scripts/gen_points.py ile PDF'lerden otomatik çıkarılan noktalar).
Üretimden sonra Excel dosyası kullanıcının ana kaynağıdır; betiği yeniden çalıştırmak
Excel'de yapılan elle düzenlemelerin üzerine yazar (önce dosya tarihini/kilidini kontrol edin).

Kullanım:  python -X utf8 scripts/build_database.py
"""
import os
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter as CL

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.environ.get("COAX_OUT") or os.path.join(ROOT, "TA3HX_Coax_Database.xlsx")
LEGACY = os.path.join(ROOT, "data", "legacy", "TA3HRJ - Coaxial_Cable_Attenuation_Chart.xlsx")
ACCESSED = "2026-10-08"
ACCESSED_OVERRIDE = {"S012": "2026-10-09", "S013": "2026-10-09"}

# ----------------------------------------------------------------- stiller
F = "Arial"
f_base = Font(name=F, size=10)
f_input = Font(name=F, size=10, color="0000FF")
f_calc = Font(name=F, size=10, color="000000")
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color="1F3864")
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_hdr_calc = Font(name=F, size=10, bold=True, color="000000")
fill_calc_hdr = PatternFill("solid", fgColor="D9D9D9")
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_out = PatternFill("solid", fgColor="E2EFDA")
fill_band = PatternFill("solid", fgColor="1F3864")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

FT_PER_M = 0.3048          # dB/100 m -> dB/100 ft çarpanı
M_PER_FT_100 = 100 / 30.48  # dB/100 ft -> dB/100 m çarpanı

# ===================================================================== KAYNAKLAR
SOURCES = [
    # id, yayıncı, belge, parça no, revizyon, belge tarihi, barındıran, url, yerel dosya, not
    ("S000", "TA3HX (o dönemki çağrı işareti TA3HRJ)", "50 ohm Cable - Nominal attenuation of 30.5 metres (100ft)", "-", "-", "-",
     "Kullanıcı", "", "data/legacy/TA3HRJ - Coaxial_Cable_Attenuation_Chart.xlsx",
     "Orijinal elle derlenmiş tablo. Bant aralığı sütunları (ör. 148-174 MHz) tek frekansa bağlı değil."),
    ("S001", "Times Microwave Systems", "Communications Coax Selection Guide (katalog s.202-203)", "LMR serisi", "-", "2007-2008 (PDF tarihi)",
     "Distribütör kopyası (RF Parts)", "https://www.rfparts.com/old_site/pdf_docs/LMR/Coax_Selection_Guide.pdf",
     "kaynaklar/S001_TimesMicrowave_Coax_Selection_Guide.pdf",
     "LMR k1/k2 katsayıları (dB/100ft, f MHz, +25°C) bu belgeden. Tablodaki rakip kablo verileri 'Competitor's data as published', * işaretliler tahmini."),
    ("S002", "Times Microwave Systems", "LMR-400 Low Loss Flexible Coax Cable Data Sheet", "LMR-400", "REV 1.3", "2020 (PDF 2023-09-29)",
     "Distribütör kopyası (Fairview Microwave)", "https://www.fairviewmicrowave.com/images/productPDF/LMR-400.pdf",
     "kaynaklar/S002_TimesMicrowave_LMR-400_Fairview.pdf",
     "timesmicrowave.com doğrudan erişime 403 döndü; Fairview kopyası kullanıldı."),
    ("S003", "Times Microwave Systems", "LMR-240 Low Loss Flexible Coax Cable Data Sheet", "LMR-240", "-", "2023 (PDF 2024-06-27)",
     "Distribütör kopyası (Fairview Microwave)", "https://www.fairviewmicrowave.com/images/productPDF/LMR-240.pdf",
     "kaynaklar/S003_TimesMicrowave_LMR-240_Fairview.pdf", ""),
    ("S004", "Belden", "Product: 8259 Technical Data", "8259", "0.533", "2026-02-20",
     "Üretici", "https://catalog.belden.com/techdata/EN/8259_techdata.pdf",
     "kaynaklar/S004_Belden_8259.pdf", ""),
    ("S005", "Belden", "Product: 8267 Technical Data", "8267", "0.599", "2026-02-20",
     "Üretici", "https://catalog.belden.com/techdata/EN/8267_techdata.pdf",
     "kaynaklar/S005_Belden_8267.pdf", ""),
    ("S006", "Belden", "Product: 8216 Technical Data", "8216", "0.540", "2026-02-20",
     "Üretici", "https://catalog.belden.com/techdata/EN/8216_techdata.pdf",
     "kaynaklar/S006_Belden_8216.pdf", ""),
    ("S007", "Belden", "Product: 9258 Technical Data", "9258", "0.514", "2026-02-20",
     "Üretici", "https://catalog.belden.com/techdata/EN/9258_techdata.pdf",
     "kaynaklar/S007_Belden_9258.pdf", ""),
    ("S008", "Belden", "Product: H155A01 Technical Data", "H155A01", "0.213", "2026-02-20",
     "Üretici", "https://catalog.belden.com/techdata/EN/H155A01_techdata.pdf",
     "kaynaklar/S008_Belden_H155A01.pdf", "PE kılıflı H155 (siyah). PVC kılıflı kardeşi H155A00."),
    ("S009", "Messi & Paoloni", "Hyperflex 10 - Full Datasheet ENG", "Hyperflex 10", "-", "2021-02-08 (PDF tarihi)",
     "Üretici", "https://messi.it/dati/layout/files/CartellaElementi/Hyperflex 10 - Full Datasheet ENG.pdf",
     "kaynaklar/S009_MessiPaoloni_Hyperflex10.pdf",
     "RigExpert'te barındırılan 2024 kopyasında güç değerleri farklı (1.8 MHz'de 5581 W, tepe 10 kW); hangisinin güncel olduğu doğrulanamadı. "
     "8000/9000 MHz değerleri PDF metninde iç içe geçmiş; 44.2 ve 47.5 dB/100m olarak çözümlendi."),
    ("S010", "Messi & Paoloni", "Airborne 5 - Full Datasheet ENG", "Airborne 5", "-", "2026-03-03 (PDF tarihi)",
     "Üretici", "https://messi.it/dati/layout/files/CartellaElementi/Airborne 5 - Full Datasheet ENG.pdf",
     "kaynaklar/S010_MessiPaoloni_Airborne5.pdf",
     "Örgü (braid) kaplama oranı PDF metninden okunamadı."),
    ("S011", "CommScope", "LDF4-50A HELIAX Low Density Foam Coaxial Cable - Product Specifications", "LDF4-50A", "Revised June 12, 2020", "2020-06-12",
     "Distribütör kopyası (Tempest NS)", "https://www.tempestns.com/wp-content/uploads/2020/06/Commscope-LDF4-50A-data.pdf",
     "kaynaklar/S011_CommScope_LDF4-50A.pdf",
     "commscope.com üzerindeki 115611 numaralı resmî bağlantı 404 döndü."),
    ("S012", "Konnektör Sepeti (satıcı)", "N erkek - N erkek Koaksiyel Kablo LMR400/RWC400 - ürün sayfası", "RWC400PE (stok kodu AUNN-155-RWC400)",
     "-", "-", "Satıcı ürün sayfası", "https://www.konnektorsepeti.com/hazir-kablo/n-n-kablo/n-erkek-n-erkek-koaksiyel-kablo-lmr400-rwc400",
     "kaynaklar/S012_KonnektorSepeti_RWC400PE.txt",
     "Hazır N-N kablo ürün sayfası. Kablo üreticisi sayfada yazmıyor; teknik metin Reçber RWC 400 katalog sayfasıyla (S013) neredeyse birebir aynı. "
     "Farklar: empedans toleransı ±2 Ω (katalog ±3 Ω), korozif gaz standardı EN 50267-2-1/2 (katalog EN 60754-1/2). Sayfa metni 2026-10-09 tarihinde arşivlendi."),
    ("S013", "Reçber Kablo", "Reçber Kablo Ürün Kataloğu - Koaksiyel Kablolar, RWC 400 (katalog s.71)", "RWC 400 PE (ürün kodu 305064)",
     "-", "2017-04-10 (PDF tarihi)", "Distribütör kopyası (Trimegs)",
     "https://www.trimegs.com/image/catalog/catalog-documents/recber-kablo-urun-katalogu.pdf",
     "kaynaklar/S013_Recber_Kablo_Urun_Katalogu.pdf",
     "222 sayfalık katalog; RWC 200 (s.69), RWC 240 (s.70), RWC 400 (s.71), RG 213 U (s.68). Zayıflama değerleri MAKSİMUM (20 °C). "
     "Kullanıcı satıcı stok kodundan (AUNN-155-RWC400) bu kataloğa ulaştı."),
]

# ===================================================================== KABLOLAR
# Sıra: kayıp sırasına göre (ince -> kalın). Bu sıra Ham_Bands raporuna da yansır.
CABLE_COLS = [
    # key, Türkçe açıklama, birim
    ("cable_id", "Benzersiz kablo kimliği (ÜRETİCİ-PARÇA)", ""),
    ("display_name", "Görünen ad", ""),
    ("manufacturer", "Üretici", ""),
    ("part_number", "Üretici parça numarası", ""),
    ("family", "Tip / aile (RG eşdeğeri vb.)", ""),
    ("short_name", "Kısa ad (web kartı / seçici)", ""),
    ("size_class", "Çap sınıfı: thin, rg58, rg8x, c7, rg213, large (od_mm'den)", ""),
    ("popular_rank", "Sık önerilenler sırası (1-8 öne çıkan; boş = listede değil)", ""),
    ("standard_type", "Sınıfının standart RG tipi mi (Evet / boş)", ""),
    ("value_type", "Zayıflama değer türü: typ = nominal/tipik, max = maksimum", ""),
    ("impedance_ohm", "Karakteristik empedans", "Ω"),
    ("impedance_tol_ohm", "Empedans toleransı (±)", "Ω"),
    ("od_mm", "Dış çap (kılıf üzeri)", "mm"),
    ("inner_material", "İç iletken malzemesi", ""),
    ("inner_construction", "İç iletken yapısı", ""),
    ("inner_od_mm", "İç iletken çapı", "mm"),
    ("dielectric_material", "Dielektrik malzemesi", ""),
    ("dielectric_od_mm", "Dielektrik çapı", "mm"),
    ("shield", "Ekran yapısı", ""),
    ("braid_coverage_pct", "Örgü kaplama oranı", "%"),
    ("jacket", "Dış kılıf", ""),
    ("velocity_factor", "Hız faktörü (VF)", "-"),
    ("capacitance_pf_m", "Kapasitans", "pF/m"),
    ("inductance_uh_m", "Endüktans", "µH/m"),
    ("dcr_inner_ohm_km", "İç iletken DC direnci", "Ω/km"),
    ("dcr_outer_ohm_km", "Dış iletken (ekran) DC direnci", "Ω/km"),
    ("shielding_db", "Ekranlama etkinliği (min.)", "dB"),
    ("fmax_mhz", "Üreticinin belirttiği maks. frekans", "MHz"),
    ("peak_power_kw", "Tepe güç", "kW"),
    ("voltage_v", "Gerilim değeri", "V"),
    ("voltage_note", "Gerilim değerinin türü", ""),
    ("bend_radius_single_mm", "Min. bükülme yarıçapı - tek/kurulum", "mm"),
    ("bend_radius_repeated_mm", "Min. bükülme yarıçapı - tekrarlı", "mm"),
    ("weight_kg_km", "Ağırlık", "kg/km"),
    ("tensile_kg", "Maks. çekme kuvveti", "kg"),
    ("temp_min_c", "Çalışma sıcaklığı min.", "°C"),
    ("temp_max_c", "Çalışma sıcaklığı maks.", "°C"),
    ("outdoor_uv", "Dış mekan / UV dayanımı", ""),
    ("direct_burial", "Toprağa doğrudan gömme", ""),
    ("fire_class", "Yangın sınıfı / onaylar", ""),
    ("connectors", "Üreticinin listelediği konnektör aileleri", ""),
    ("k1_pub_db_100ft", "Üreticinin yayımladığı k1 (α=k1·√f+k2·f, dB/100ft)", "dB/100ft"),
    ("k2_pub_db_100ft", "Üreticinin yayımladığı k2", "dB/100ft"),
    ("primary_source_id", "Ana kaynak", ""),
    ("notes", "Notlar", ""),
]
CABLE_CALC_COLS = [
    ("att_first_row", "Attenuation sayfasında ilk nokta satırı (formül)", ""),
    ("att_n", "Zayıflama noktası sayısı (formül)", ""),
    ("f_min_data", "Datasheet verisinin en düşük frekansı (formül)", "MHz"),
    ("f_max_data", "Datasheet verisinin en yüksek frekansı (formül)", "MHz"),
    ("k0_fit", "QA modeli sabit terim (α=k0+k1·√f+k2·f, dB/100m)", "dB/100m"),
    ("k1_fit", "QA modeli k1", "dB/100m"),
    ("k2_fit", "QA modeli k2", "dB/100m"),
    ("fit_rms_pct", "QA modeli ile datasheet arası RMS sapma", "%"),
    ("pwr_first_row", "Power sayfasında ilk nokta satırı (formül)", ""),
    ("pwr_n", "Güç noktası sayısı (formül)", ""),
]

CABLES = [
    dict(cable_id="BEL-8216", display_name="Belden 8216 (RG-174)", manufacturer="Belden", part_number="8216", family="RG-174",
         impedance_ohm=50, od_mm=2.79, inner_material="Bakır kaplı çelik (BCCS)", inner_construction="26 AWG, 7x34 örgülü",
         inner_od_mm=0.48, dielectric_material="PE (dolu)", dielectric_od_mm=1.5, shield="Kalaylı bakır örgü", braid_coverage_pct=90,
         jacket="PVC", velocity_factor=0.66, capacitance_pf_m=101, dcr_inner_ohm_km=318.2, dcr_outer_ohm_km=35.1,
         voltage_v=1100, voltage_note="Gerilim değeri (UL AWM 1354: 30 V)", bend_radius_single_mm=28, weight_kg_km=11.9, tensile_kg=9.1,
         temp_min_c=-30, temp_max_c=75, outdoor_uv="Dış mekan - havai (sadece siyah, taşıyıcı tel ile)", fire_class="CPR Eca",
         primary_source_id="S006",
         notes="İç iletken çelik çekirdekli: düşük frekansta kayıp √f yasasından sapar (1 MHz'de 1.9 dB/100ft)."),
    dict(cable_id="MP-AIRBORNE5", display_name="Messi & Paoloni Airborne 5", manufacturer="Messi & Paoloni", part_number="Airborne 5",
         family="5 mm düşük kayıplı (RG-58 boyutu)", impedance_ohm=50, impedance_tol_ohm=3, od_mm=5.0, inner_material="Bakır (%99.9, çıplak)",
         inner_construction="Tek tel", inner_od_mm=1.13, dielectric_material="Fiziksel köpük PE (3 katlı)", dielectric_od_mm=3.0,
         shield="Al-polyester-Al folyo (%100) + örgü", jacket="UV korumalı PE", velocity_factor=0.85, capacitance_pf_m=76,
         dcr_inner_ohm_km=17, dcr_outer_ohm_km=34, shielding_db=105, peak_power_kw=2.0,
         bend_radius_single_mm=25, bend_radius_repeated_mm=50, weight_kg_km=23, temp_min_c=-45, temp_max_c=70,
         outdoor_uv="Evet", direct_burial="Evet", fire_class="CPR Fca (DoP MP0095)", primary_source_id="S010",
         notes="Ekranlama >105 dB (100-2000 MHz). Tekrarlı bükülme: 15 bükülmeye kadar."),
    dict(cable_id="BEL-8259", display_name="Belden 8259 (RG-58)", manufacturer="Belden", part_number="8259", family="RG-58",
         impedance_ohm=50, od_mm=4.90, inner_material="Kalaylı bakır", inner_construction="20 AWG, 19x33 örgülü", inner_od_mm=0.89,
         dielectric_material="PE (dolu)", dielectric_od_mm=2.97, shield="Kalaylı bakır örgü", braid_coverage_pct=95, jacket="PVC",
         velocity_factor=0.66, capacitance_pf_m=101, dcr_inner_ohm_km=32, dcr_outer_ohm_km=12, voltage_v=1900,
         voltage_note="Gerilim değeri", bend_radius_single_mm=48, weight_kg_km=37.2, tensile_kg=19, temp_min_c=-30, temp_max_c=75,
         outdoor_uv="İç mekan; havai (sadece siyah, taşıyıcı tel ile)", fire_class="CPR Eca", primary_source_id="S004",
         notes="Datasheet güç tablosu vermiyor."),
    dict(cable_id="BEL-H155A01", display_name="Belden H155A01 (H155 PE)", manufacturer="Belden", part_number="H155A01", family="H155",
         impedance_ohm=50, od_mm=5.4, inner_material="Çıplak bakır", inner_construction="16 AWG örgülü", inner_od_mm=1.41,
         dielectric_material="Köpük PE", dielectric_od_mm=3.9, shield="Al/poly/Al folyo (%100) + kalaylı bakır örgü", braid_coverage_pct=80,
         jacket="PE (siyah)", velocity_factor=0.80, capacitance_pf_m=84, dcr_inner_ohm_km=15.4, shielding_db=85,
         tensile_kg=10.2, temp_min_c=-30, temp_max_c=70, outdoor_uv="Dış mekan", fire_class="Halojensiz (IEC 60754)",
         primary_source_id="S008",
         notes="DCR değeri maksimumdur. Ekranlama 30-1000 MHz için. Maks. zayıflama nominalin %10 üstü. Datasheet güç tablosu vermiyor."),
    dict(cable_id="BEL-9258", display_name="Belden 9258 (RG-8X)", manufacturer="Belden", part_number="9258", family="RG-8X (mini RG-8)",
         impedance_ohm=50, od_mm=6.15, inner_material="Çıplak bakır", inner_construction="16 AWG, 19x29 örgülü", inner_od_mm=1.45,
         dielectric_material="Köpük PE", dielectric_od_mm=3.94, shield="Çıplak bakır örgü", braid_coverage_pct=95, jacket="PVC",
         velocity_factor=0.82, capacitance_pf_m=83, dcr_inner_ohm_km=13.45, dcr_outer_ohm_km=10, voltage_v=300,
         voltage_note="UL CM değeri", bend_radius_single_mm=61, weight_kg_km=55.1, tensile_kg=28, temp_min_c=-40, temp_max_c=80,
         outdoor_uv="İç/dış mekan", fire_class="UL CM, CPR Eca", primary_source_id="S007", notes=""),
    dict(cable_id="TMS-LMR240", display_name="Times Microwave LMR-240", manufacturer="Times Microwave Systems", part_number="LMR-240",
         family="LMR (RG-8X alternatifi)", impedance_ohm=50, od_mm=6.1, inner_material="Bakır", inner_construction="Tek tel",
         inner_od_mm=1.42, dielectric_material="Köpük PE", dielectric_od_mm=3.81, shield="Al bant (yapışık) + kalaylı bakır örgü",
         jacket="PE (siyah, UV korumalı)", velocity_factor=0.84, capacitance_pf_m=79.4, inductance_uh_m=0.2,
         dcr_inner_ohm_km=10.50, dcr_outer_ohm_km=12.76, shielding_db=90, fmax_mhz=8000, peak_power_kw=5.6, voltage_v=1500,
         voltage_note="Dielektrik dayanım (DC)", bend_radius_single_mm=19.05, bend_radius_repeated_mm=63.5, weight_kg_km=49.1,
         tensile_kg=36.29, temp_min_c=-40, temp_max_c=85, outdoor_uv="Evet", k1_pub_db_100ft=0.24208, k2_pub_db_100ft=0.00033,
         primary_source_id="S003", notes="k1/k2: S001. Kesim (cutoff) frekansı 31 GHz."),
    dict(cable_id="BEL-8267", display_name="Belden 8267 (RG-213/U)", manufacturer="Belden", part_number="8267", family="RG-213/U",
         impedance_ohm=50, od_mm=10.3, inner_material="Çıplak bakır", inner_construction="13 AWG, 7x21 örgülü", inner_od_mm=2.26,
         dielectric_material="PE (dolu)", dielectric_od_mm=7.24, shield="Çıplak bakır örgü", braid_coverage_pct=95,
         jacket="PVC (kirletmeyen)", velocity_factor=0.66, capacitance_pf_m=101, dcr_inner_ohm_km=5.58, dcr_outer_ohm_km=3.9,
         voltage_v=3700, voltage_note="Mil-Spec değeri (UL CMX: 300 V)", bend_radius_single_mm=100, weight_kg_km=154.8,
         tensile_kg=83.5, temp_min_c=-40, temp_max_c=80, outdoor_uv="İç mekan", fire_class="UL CMX, VW-1, CPR Eca",
         primary_source_id="S005", notes="Ticari, QPL dışı (M17/163-00001 karşılığı)."),
    dict(cable_id="MP-HYPERFLEX10", display_name="Messi & Paoloni Hyperflex 10", manufacturer="Messi & Paoloni", part_number="Hyperflex 10",
         family="10.3 mm düşük kayıplı (RG-213 boyutu)", impedance_ohm=50, impedance_tol_ohm=3, od_mm=10.3, inner_material="Bakır",
         inner_construction="19x0.59 mm örgülü", inner_od_mm=2.9, dielectric_material="Fiziksel köpük PE (3 katlı)", dielectric_od_mm=7.3,
         shield="Bakır folyo (%100) + bakır kaplı alüminyum örgü (192 tel)", braid_coverage_pct=85, jacket="UV dayanımlı PVC (siyah)",
         velocity_factor=0.87, capacitance_pf_m=78, dcr_inner_ohm_km=3.6, dcr_outer_ohm_km=12, shielding_db=105, peak_power_kw=13,
         bend_radius_single_mm=40, bend_radius_repeated_mm=80, weight_kg_km=112, temp_min_c=-40, temp_max_c=60, outdoor_uv="Evet (UV dayanımlı)",
         fire_class="IEC 60332-1-2, CPR (DoP MP00103)", connectors="UHF (PL-259), N, BNC, SMA, TNC, 7/16",
         primary_source_id="S009", notes="Kılıf kıvılcım testi 8 kV. Tekrarlı bükülme: 15 bükülmeye kadar."),
    dict(cable_id="RCB-RWC400PE", display_name="Reçber RWC 400 PE (LMR-400 muadili)", manufacturer="Reçber Kablo",
         part_number="RWC 400 PE (305064)", family="LMR-400 muadili (10.2 mm)", impedance_ohm=50, impedance_tol_ohm=3, od_mm=10.2,
         inner_material="Elektrolitik bakır", inner_construction="Tek tel (Class 1, IEC 60228)", inner_od_mm=2.70,
         dielectric_material="Fiziksel köpük PE (3 katlı)", dielectric_od_mm=7.20, shield="Al-Pet-Al folyo (%100) + kalaylı bakır örgü",
         braid_coverage_pct=85, jacket="PE siyah RAL 9011 (80 °C, EN 50290-2-24); PVC (305057) ve HFFR (305071) seçenekleri var",
         velocity_factor=0.85, capacitance_pf_m=78.5, shielding_db=90, voltage_v=2000, voltage_note="Maks. çalışma gerilimi (test gerilimi 5000 V)",
         bend_radius_single_mm=102, weight_kg_km=128, temp_min_c=-30, temp_max_c=70, outdoor_uv="Evet (PE kılıf: açık hava)",
         direct_burial="Üreticiye göre PE kılıf yeraltı uygulamalarına uygun",
         fire_class="EN 60332-1-2 (PVC/HFFR); HFFR: EN 60754-1/2, EN 61034-2; segregasyon sınıfı d (EN 50174-2)",
         connectors="BNC, TNC, SMA, N (üretici metni)", primary_source_id="S013",
         notes="ZAYIFLAMA DEĞERLERİ MAKSİMUMDUR (diğer kabloların çoğu tipik/nominal) - karşılaştırmada dikkat. Bükülme yarıçapı 10×D. "
               "VF %85±2, kapasite ±2 pF/m. Yalıtım direnci min. 2 GΩ·km. EN 50117, IEC 61196. Bakır ağırlığı 82 kg/km. "
               "Geri dönüş kaybı >23 dB (30-1000 MHz), >20 dB (1-3 GHz). Güç değeri yayımlanmamış. "
               "Konnektör Sepeti AUNN-155-RWC400 hazır N-N kablosu bu kabloyu kullanıyor (S012; o sayfa empedansı ±2 Ω yazıyor)."),
    dict(cable_id="TMS-LMR400", display_name="Times Microwave LMR-400", manufacturer="Times Microwave Systems", part_number="LMR-400",
         family="LMR (RG-8/RG-213 alternatifi)", impedance_ohm=50, od_mm=10.29, inner_material="Bakır kaplı alüminyum (CCA)",
         inner_construction="Tek tel", inner_od_mm=2.74, dielectric_material="Köpük PE", dielectric_od_mm=7.24,
         shield="Al bant (yapışık) + kalaylı bakır örgü", jacket="PE (siyah, UV korumalı)", velocity_factor=0.85,
         capacitance_pf_m=78.41, inductance_uh_m=0.2, dcr_inner_ohm_km=4.56, dcr_outer_ohm_km=5.41, shielding_db=90, fmax_mhz=6000,
         peak_power_kw=16, voltage_v=2500, voltage_note="Dielektrik dayanım (DC)", bend_radius_single_mm=25.4,
         bend_radius_repeated_mm=101.6, weight_kg_km=99.7, tensile_kg=72.57, temp_min_c=-40, temp_max_c=85, outdoor_uv="Evet",
         k1_pub_db_100ft=0.12229, k2_pub_db_100ft=0.00026, primary_source_id="S002",
         notes="k1/k2: S001. Fairview kopyası frekans aralığını DC-6 GHz verir; tabloda 5.8 GHz değeri var."),
    dict(cable_id="CS-LDF4-50A", display_name="CommScope LDF4-50A (1/2\" HELIAX)", manufacturer="CommScope", part_number="LDF4-50A",
         family="1/2\" HELIAX LDF (oluklu)", impedance_ohm=50, impedance_tol_ohm=1, od_mm=15.875, inner_material="Bakır kaplı alüminyum (CCA)",
         inner_construction="Tek tel", inner_od_mm=4.826, dielectric_material="Köpük PE", dielectric_od_mm=12.954,
         shield="Oluklu bakır boru (13.97 mm)", jacket="PE (siyah, halojensiz)", velocity_factor=0.88, capacitance_pf_m=75.8,
         inductance_uh_m=0.19, dcr_inner_ohm_km=1.48, dcr_outer_ohm_km=2.69, fmax_mhz=8800, peak_power_kw=40, voltage_v=4000,
         voltage_note="DC test gerilimi", bend_radius_single_mm=50.8, bend_radius_repeated_mm=127, weight_kg_km=220, tensile_kg=113,
         temp_min_c=-55, temp_max_c=85, outdoor_uv="Evet", fire_class="CPR Fca", primary_source_id="S011",
         notes="Zayıflama değerleri tipik, %5 içinde garanti. Kurulum sıcaklığı -40..+60 °C."),
]

# ===================================================================== ZAYIFLAMA NOKTALARI
# (cable_id, source_id, unit, temp_c, source_ref, [(f MHz, değer), ...], not)
ATT = [
    ("BEL-8216", "S006", "dB/100ft", None, "Attenuation tablosu (Nom.)",
     [(1, 1.9), (10, 3.3), (50, 5.8), (100, 8.4), (200, 12.5), (400, 19.0), (700, 27.0), (900, 31.0), (1000, 34.0)], ""),
    ("MP-AIRBORNE5", "S010", "dB/100m", 20, "ATTENUATION (20°C) tablosu",
     [(1.8, 1.7), (3.5, 2.3), (7, 3.0), (10, 3.4), (14, 4.0), (21, 4.8), (28, 5.5), (50, 7.1), (100, 9.4), (144, 11.1), (200, 12.8),
      (400, 18.3), (430, 19.0), (800, 26.5), (1000, 29.8), (1296, 34.2), (2400, 47.5), (3000, 53.5), (4000, 61.0), (5000, 68.6),
      (6000, 75.6)], ""),
    ("BEL-8259", "S004", "dB/100ft", None, "Attenuation tablosu (Nom.)",
     [(1, 0.4), (10, 1.5), (50, 3.7), (100, 5.4), (200, 8.1), (400, 12.4), (700, 17.7), (900, 21.1), (1000, 22.8)], ""),
    ("BEL-H155A01", "S008", "dB/100m", None, "Attenuation tablosu (Nom.)",
     [(5, 2.5), (50, 6.9), (100, 9.1), (230, 13.4), (400, 18.0), (800, 26.1), (862, 27.3), (1000, 29.6), (1350, 34.9), (1750, 40.3),
      (2150, 46.0), (2400, 49.1), (3000, 56.3), (3600, 62.9), (4200, 69.1), (4800, 75.1), (5400, 80.8), (6000, 86.5)], ""),
    ("BEL-9258", "S007", "dB/100ft", None, "Attenuation tablosu (Nom.)",
     [(1, 0.3), (10, 0.9), (50, 2.1), (100, 3.1), (200, 4.5), (400, 6.6), (700, 9.1), (900, 10.7), (1000, 11.2)], ""),
    ("TMS-LMR240", "S001", "dB/100ft", 25, "Attenuation tablosu, LMR-240 sütunu",
     [(30, 1.3), (700, 6.6)], ""),
    ("TMS-LMR240", "S003", "dB/100ft", None, "Performance by Frequency Band (Typ)",
     [(50, 1.7), (150, 3.0), (220, 3.7), (450, 5.3), (900, 7.6), (1500, 9.9), (1800, 10.9), (2000, 11.5), (2500, 12.9), (5800, 20.4)], ""),
    ("BEL-8267", "S005", "dB/100ft", None, "Attenuation tablosu (Nom.)",
     [(1, 0.17), (10, 0.55), (50, 1.3), (100, 1.9), (200, 2.7), (400, 4.1), (700, 6.5), (900, 7.6), (1000, 8.0), (4000, 21.5)], ""),
    ("MP-HYPERFLEX10", "S009", "dB/100m", 20, "ATTENUATION (20°C) tablosu",
     [(1.8, 0.8), (3.5, 1.0), (7, 1.1), (10, 1.3), (14, 1.5), (21, 1.8), (28, 2.0), (50, 2.7), (100, 3.9), (144, 4.7), (200, 5.6),
      (400, 8.3), (430, 8.6), (800, 11.9), (1000, 13.4), (1296, 15.4), (2400, 21.8), (3000, 24.6), (4000, 29.1), (5000, 33.1),
      (6000, 36.9), (7000, 40.7), (8000, 44.2), (9000, 47.5), (10000, 50.7)], ""),
    ("RCB-RWC400PE", "S013", "dB/100m", 20, "Katalog s.71, Zayıflama @20 °C max. (S012 satıcı sayfasıyla aynı)",
     [(30, 2.38), (50, 3.01), (150, 5.10), (220, 6.18), (450, 8.96), (900, 13.02), (1500, 17.33), (1800, 19.24), (2000, 20.45),
      (2500, 23.30), (5800, 38.92)], "Maksimum değer"),
    ("TMS-LMR400", "S001", "dB/100ft", 25, "Attenuation tablosu, LMR-400 sütunu",
     [(30, 0.7), (700, 3.42)], ""),
    ("TMS-LMR400", "S002", "dB/100ft", None, "Performance by Frequency Band (Typ)",
     [(50, 0.9), (150, 1.5), (220, 1.9), (450, 2.7), (900, 3.9), (1500, 5.1), (1800, 5.7), (2000, 6.0), (2500, 6.8), (5800, 10.8)], ""),
]
LDF4 = [(1.0, 0.211, 36.11), (1.5, 0.259, 29.46), (2.0, 0.299, 25.5), (10, 0.672, 11.35), (20, 0.954, 7.99), (30, 1.172, 6.51),
        (50, 1.521, 5.02), (85, 1.995, 3.82), (88, 2.031, 3.76), (100, 2.169, 3.52), (108, 2.256, 3.38), (150, 2.673, 2.85),
        (174, 2.887, 2.64), (200, 3.103, 2.46), (204, 3.135, 2.43), (300, 3.835, 1.99), (400, 4.462, 1.71), (450, 4.749, 1.61),
        (460, 4.804, 1.59), (500, 5.021, 1.52), (512, 5.085, 1.5), (600, 5.533, 1.38), (700, 6.009, 1.27), (800, 6.456, 1.18),
        (824, 6.56, 1.16), (894, 6.855, 1.11), (960, 7.124, 1.07), (1000, 7.284, 1.05), (1218, 8.11, 0.94), (1250, 8.226, 0.93),
        (1500, 9.093, 0.84), (1700, 9.744, 0.78), (1794, 10.039, 0.76), (1800, 10.058, 0.76), (2000, 10.666, 0.72),
        (2100, 10.961, 0.7), (2200, 11.251, 0.68), (2300, 11.535, 0.66), (2500, 12.09, 0.63), (2700, 12.627, 0.6),
        (3000, 13.407, 0.57), (3400, 14.401, 0.53), (3600, 14.882, 0.51), (3700, 15.118, 0.5), (3800, 15.353, 0.5),
        (3900, 15.585, 0.49), (4000, 15.815, 0.48), (4100, 16.042, 0.48), (4200, 16.268, 0.47), (4300, 16.492, 0.46),
        (4400, 16.714, 0.46), (4500, 16.934, 0.45), (4600, 17.153, 0.44), (4700, 17.37, 0.44), (4800, 17.585, 0.43),
        (4900, 17.798, 0.43), (5000, 18.01, 0.42), (6000, 20.055, 0.38), (8000, 23.826, 0.32), (8800, 25.244, 0.3)]
ATT.append(("CS-LDF4-50A", "S011", "dB/100m", 20, "Attenuation tablosu", [(f, a) for f, a, _ in LDF4], ""))

# ===================================================================== GÜÇ NOKTALARI
PWR = [
    ("BEL-8216", "S006", "W", None, "Power Rating tablosu",
     [(1, 976), (10, 354), (50, 197), (100, 161), (200, 136), (400, 120), (700, 112), (900, 109), (1000, 108)]),
    ("MP-AIRBORNE5", "S010", "W", 40, "POWER HANDLING (40°C)",
     [(1.8, 1172), (3.5, 837), (7, 625), (10, 543), (14, 471), (21, 394), (28, 346), (50, 268), (100, 198), (144, 170), (200, 146),
      (400, 102), (430, 99), (800, 71), (1000, 63), (1296, 55), (2400, 39), (3000, 35), (4000, 31), (5000, 27), (6000, 25)]),
    ("BEL-9258", "S007", "W", None, "Power Rating tablosu",
     [(10, 1000), (50, 370), (100, 250), (200, 190), (400, 110), (700, 75), (1000, 60)]),
    ("TMS-LMR240", "S003", "W", None, "Input Power (CW), Max",
     [(50, 1150), (150, 660), (220, 540), (450, 380), (900, 260), (1500, 200), (1800, 180), (2000, 170), (2500, 150), (5800, 100)]),
    ("BEL-8267", "S005", "W", None, "Power Rating tablosu",
     [(1, 9295), (10, 2761), (50, 1122), (100, 748), (200, 494), (400, 326), (700, 235), (900, 205), (1000, 194), (4000, 110)]),
    ("MP-HYPERFLEX10", "S009", "W", 40, "POWER HANDLING (40°C)",
     [(1.8, 9927), (3.5, 7721), (7, 5990), (10, 5186), (14, 4483), (21, 3777), (28, 3357), (50, 2518), (100, 1759), (144, 1460),
      (200, 1226), (400, 837), (430, 808), (800, 581), (1000, 516), (1296, 449), (2400, 319), (3000, 282), (4000, 239),
      (5000, 210), (6000, 188), (7000, 171), (8000, 157), (10000, 137)]),
    ("TMS-LMR400", "S002", "W", None, "Power In (CW), Max",
     [(50, 2570), (150, 1470), (220, 1200), (450, 830), (900, 580), (1500, 440), (1800, 400), (2000, 370), (2500, 330), (5800, 210)]),
    ("CS-LDF4-50A", "S011", "kW", 40, "Attenuation tablosu - Average Power (40°C ortam)", [(f, p) for f, _, p in LDF4]),
]

# ===================================================================== AMATÖR BANTLARI
BANDS = [("160 m", 1.8), ("80 m", 3.5), ("40 m", 7), ("20 m", 14), ("15 m", 21), ("10 m", 28), ("6 m", 50),
         ("2 m", 144), ("70 cm", 432), ("23 cm", 1296), ("13 cm", 2400), ("6 cm", 5760)]

# ===================================================================== ESKİ TABLO EŞLEMESİ
LEGACY_BANDS = [("E", "1-10 MHz", 1, 10), ("F", "30-50 MHz", 30, 50), ("G", "70-100 MHz", 70, 100), ("H", "148-174 MHz", 148, 174),
                ("I", "220 MHz", 220, 220), ("J", "400-520 MHz", 400, 520), ("K", "700-960 MHz", 700, 960),
                ("L", "1.2 GHz", 1200, 1300), ("M", "2.4 GHz", 2400, 2500), ("N", "5.8 GHz", 5725, 5875)]
# Sayfa1 satırı -> (cable_id, eşleşme türü, not)
LEGACY_MAP = {
    6: ("BEL-8216", "Eşdeğer tip", "Belden 8216 RG-174 tipi."),
    9: ("BEL-8259", "Eşdeğer tip", "Belden 8259 RG-58 tipi; RG-58C/U MIL tanımından farklı olabilir."),
    15: ("BEL-9258", "Birebir", ""),
    16: ("BEL-9258", "Eşdeğer tip", "960 MHz (8.8) -> 1.2 GHz (15.9) geçişi fiziksel olarak tutarsız (%25 frekans artışına %80 kayıp artışı)."),
    19: ("BEL-H155A01", "Birebir", "PE kılıflı sürüm ile karşılaştırıldı."),
    20: ("MP-AIRBORNE5", "Birebir", ""),
    22: ("TMS-LMR240", "Birebir", ""),
    23: ("BEL-8267", "Eşdeğer tip", "Satır 25/26 (RG213/U, RG214/U) ile tekrar."),
    25: ("BEL-8267", "Eşdeğer tip", ""),
    38: ("TMS-LMR400", "Birebir", ""),
    40: ("MP-HYPERFLEX10", "Birebir", ""),
    50: ("CS-LDF4-50A", "Eşdeğer tip", "LDF4-50A zaten 1/2\" HELIAX; satır 51 ile aynı kablo. Eski çap 16 mm, datasheet 15.875 mm."),
    51: ("CS-LDF4-50A", "Birebir", "Eski çap 13.6 mm; datasheet dış çap 15.875 mm (13.97 mm oluklu bakır dış iletken)."),
}
LEGACY_NOTES = {
    21: "RG-62A/U standart olarak 93 Ω'dur; tabloda 50 Ω yazıyor - doğrulanmalı.",
    26: "RG-214/U çift gümüş kaplı örgülüdür; RG-213 ile aynı kayıp kabul edilmemeli.",
    32: "Satır 35 ile aynı ad, farklı çap (10.16 / 10.9 mm) ve farklı değerler - tekrar.",
    35: "Satır 32 ile aynı ad - tekrar.",
}

# ===================================================================== GENİŞLETME (pilot sonrası)
import sys  # noqa: E402
sys.path.insert(0, HERE)
from cables_extra import SOURCES_X, CABLES_X, ATT_MANUAL, PWR_MANUAL, LEGACY_MAP_X, LEGACY_NOTES_X, LEGACY_MAP_S2, web_fields  # noqa: E402
from points_extra import ATT_PDF, PWR_PDF  # noqa: E402

SOURCES += SOURCES_X
CABLES += CABLES_X
ATT += ATT_PDF + ATT_MANUAL
PWR += PWR_PDF + PWR_MANUAL
LEGACY_MAP.update(LEGACY_MAP_X)
for _r, _n in LEGACY_NOTES_X.items():
    LEGACY_NOTES[_r] = (LEGACY_NOTES.get(_r, "") + " " + _n).strip()
LEGACY_NO_SOURCE = set(LEGACY_NOTES_X)
for _s in SOURCES_X:
    ACCESSED_OVERRIDE[_s[0]] = "2026-10-09"

for _c in CABLES:
    _c.update(web_fields(_c))

# bütünlük kontrolleri
_ids = [c["cable_id"] for c in CABLES]
assert len(_ids) == len(set(_ids)), "tekrarlanan cable_id"
_sids = [s[0] for s in SOURCES]
assert len(_sids) == len(set(_sids)), "tekrarlanan source_id"
for _row in ATT + PWR:
    assert _row[0] in _ids, ("noktası olan ama Cables'da olmayan kablo", _row[0])
    assert _row[1] in _sids, ("tanımsız kaynak", _row[1])
for _c in CABLES:
    assert _c["primary_source_id"] in _sids, _c["cable_id"]
    assert any(a[0] == _c["cable_id"] for a in ATT), ("zayıflama noktası yok", _c["cable_id"])
for _m in list(LEGACY_MAP.values()) + list(LEGACY_MAP_S2.values()):
    assert _m[0] in _ids, _m


def _att_at(cid, f):
    """Sıralama için Python tarafında log-log interpolasyon (dB/100 m)."""
    pts = sorted({(pf, v * (M_PER_FT_100 if u == "dB/100ft" else 1))
                  for c, s, u, t, r, pp, n in ATT if c == cid for pf, v in pp})
    xs = [math.log(p[0]) for p in pts]
    ys = [math.log(p[1]) for p in pts]
    x = math.log(f)
    i = max(0, min(len(xs) - 2, sum(1 for v in xs if v <= x) - 1))
    return math.exp(ys[i] + (ys[i + 1] - ys[i]) * (x - xs[i]) / (xs[i + 1] - xs[i]))


# Rapor sırası: önce 50 Ω sınıfı (en kayıplıdan en az kayıplıya, 450 MHz'e göre), sonra 75 Ω, 93 Ω
def _zgroup(c):
    return 0 if c["impedance_ohm"] in (50, 52, 54) else c["impedance_ohm"]


CABLES.sort(key=lambda c: (_zgroup(c), -_att_at(c["cable_id"], 450)))

# Times Microwave Systems bir Amphenol şirketidir; amatör camiada "Amphenol LMR" olarak da anılır.
for _c in CABLES:
    if _c["manufacturer"] == "Times Microwave Systems":
        _c["manufacturer"] = "Times Microwave Systems (Amphenol)"
        _c["display_name"] += " (Amphenol)"
SOURCES = [s if s[1] != "Times Microwave Systems" else (s[0], "Times Microwave Systems (Amphenol)") + s[2:] for s in SOURCES]


# =====================================================================
def interp_formula(f, r0, n, sheet, xcol, ycol):
    """log-log doğrusal interpolasyon; veri aralığı dışında uç segmentin eğimiyle ekstrapolasyon."""
    i = (f"({r0}-1+MIN(MAX(IFERROR(MATCH(LN({f}),INDEX({sheet}!${xcol}:${xcol},{r0}):"
         f"INDEX({sheet}!${xcol}:${xcol},{r0}+{n}-1),1),1),1),{n}-1))")
    return (f"EXP(FORECAST(LN({f}),INDEX({sheet}!${ycol}:${ycol},{i}):INDEX({sheet}!${ycol}:${ycol},{i}+1),"
            f"INDEX({sheet}!${xcol}:${xcol},{i}):INDEX({sheet}!${xcol}:${xcol},{i}+1)))")


def add_table(ws, name, ncols, nrows, calc_formulas=None, style="TableStyleMedium2"):
    ref = f"A1:{CL(ncols)}{nrows + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    t._initialise_columns()
    for j, col in enumerate(t.tableColumns, 1):
        col.name = str(ws.cell(1, j).value)
        if calc_formulas and col.name in calc_formulas and not os.environ.get("NO_CALC_COL"):
            col.calculatedColumnFormula = TableFormula(attr_text=calc_formulas[col.name])
    ws.add_table(t)


def write_header(ws, headers, calc_keys=()):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        if h in calc_keys:
            c.font = f_hdr_calc
            c.fill = fill_calc_hdr
        else:
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.alignment = center
    ws.freeze_panes = "B2" if ws.title in ("Cables",) else "A2"


def set_widths(ws, widths):
    for k, w in widths.items():
        ws.column_dimensions[k].width = w


# =====================================================================
def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s_readme = wb.create_sheet("README")
    s_bands = wb.create_sheet("Ham_Bands")
    s_graph = wb.create_sheet("Graph")
    s_calc = wb.create_sheet("Calculator")
    s_cab = wb.create_sheet("Cables")
    s_att = wb.create_sheet("Attenuation")
    s_pwr = wb.create_sheet("Power")
    s_src = wb.create_sheet("Sources")
    s_lix = wb.create_sheet("Legacy_Index")
    s_lpt = wb.create_sheet("Legacy_Points")
    s_dd = wb.create_sheet("Data_Dictionary")

    # ------------------------------------------------------------ Sources
    src_hdr = ["source_id", "publisher", "document_title", "part_number", "revision", "document_date", "host",
               "url", "accessed", "local_file", "notes"]
    write_header(s_src, src_hdr)
    for i, s in enumerate(SOURCES, 2):
        vals = list(s[:7]) + [s[7], ACCESSED_OVERRIDE.get(s[0], ACCESSED), s[8], s[9]]
        for j, v in enumerate(vals, 1):
            c = s_src.cell(i, j, v)
            c.font = f_input
            c.alignment = wrap
        if s[7]:
            s_src.cell(i, 8).hyperlink = s[7]
        if s[8]:
            s_src.cell(i, 10).hyperlink = s[8]
    add_table(s_src, "tblSources", len(src_hdr), len(SOURCES))
    set_widths(s_src, {"A": 10, "B": 20, "C": 42, "D": 14, "E": 14, "F": 18, "G": 24, "H": 50, "I": 11, "J": 40, "K": 70})

    # ------------------------------------------------------------ Attenuation
    att_hdr = ["point_id", "cable_id", "freq_mhz", "value", "unit", "temp_c", "source_id", "source_ref",
               "att_db_100m", "ln_f", "ln_att", "w0", "w1", "w2", "one", "model_db_100m", "dev_pct", "dev_sq", "order_ok", "note"]
    att_calc = ["att_db_100m", "ln_f", "ln_att", "w0", "w1", "w2", "one", "model_db_100m", "dev_pct", "dev_sq", "order_ok"]
    write_header(s_att, att_hdr, att_calc)
    # Cables sütun harfleri (QA katsayıları)
    cab_keys = [c[0] for c in CABLE_COLS] + [c[0] for c in CABLE_CALC_COLS]
    CC = {k: CL(i + 1) for i, k in enumerate(cab_keys)}

    def att_formulas(r):
        k = lambda key: f"INDEX(Cables!${CC[key]}:${CC[key]},MATCH(B{r},Cables!$A:$A,0))"
        return {
            "att_db_100m": f'=IF(E{r}="dB/100ft",D{r}*100/30.48,D{r})',
            "ln_f": f"=LN(C{r})",
            "ln_att": f"=LN(I{r})",
            "w0": f"=1/I{r}",
            "w1": f"=SQRT(C{r})/I{r}",
            "w2": f"=C{r}/I{r}",
            "one": "=1",
            "model_db_100m": f"={k('k0_fit')}+{k('k1_fit')}*SQRT(C{r})+{k('k2_fit')}*C{r}",
            "dev_pct": f"=P{r}/I{r}-1",
            "dev_sq": f"=Q{r}^2",
            "order_ok": f"=IF(B{r}=B{r - 1},C{r}>C{r - 1},COUNTIF(B$1:B{r - 1},B{r})=0)",
        }

    # sırala: kablo sırası CABLES listesindeki gibi, frekans artan
    order = {c["cable_id"]: i for i, c in enumerate(CABLES)}
    rows = []
    for cid, sid, unit, temp, ref, pts, note in ATT:
        for f, v in pts:
            rows.append((cid, f, v, unit, temp, sid, ref, note))
    rows.sort(key=lambda x: (order[x[0]], x[1]))
    for i, (cid, f, v, unit, temp, sid, ref, note) in enumerate(rows, 2):
        vals = [f"A{i - 1:04d}", cid, f, v, unit, temp, sid, ref]
        for j, val in enumerate(vals, 1):
            s_att.cell(i, j, val).font = f_input
        fm = att_formulas(i)
        for j, key in enumerate(att_hdr, 1):
            if key in fm:
                c = s_att.cell(i, j, fm[key])
                c.font = f_calc
        s_att.cell(i, 20, note).font = f_input
    n_att = len(rows)
    calc2 = {k: v[1:] for k, v in att_formulas(2).items()}
    add_table(s_att, "tblAttenuation", len(att_hdr), n_att, calc2)
    for col, fmt in (("I", "0.000"), ("J", "0.000"), ("K", "0.000"), ("L", "0.0000"), ("M", "0.0000"), ("N", "0.0000"),
                     ("P", "0.000"), ("Q", "+0.0%;-0.0%;0.0%"), ("R", "0.0000")):
        for r in range(2, n_att + 2):
            s_att[f"{col}{r}"].number_format = fmt
    last = n_att + 1
    s_att.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=["ABS(Q2)>0.1"], fill=PatternFill("solid", fgColor="F8CBAD")))
    s_att.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=["ABS(Q2)>0.05"], fill=PatternFill("solid", fgColor="FFE699")))
    s_att.conditional_formatting.add(f"S2:S{last}", FormulaRule(formula=["S2=FALSE"], fill=PatternFill("solid", fgColor="FF7C80")))
    dv = DataValidation(type="list", formula1='"dB/100m,dB/100ft"', allow_blank=False)
    s_att.add_data_validation(dv)
    dv.add(f"E2:E{last + 500}")
    set_widths(s_att, {"A": 9, "B": 17, "C": 10, "D": 8, "E": 10, "F": 8, "G": 10, "H": 34, "I": 11, "J": 8, "K": 8,
                       "L": 8, "M": 8, "N": 8, "O": 5, "P": 13, "Q": 9, "R": 8, "S": 9, "T": 30})
    for col in "JKLMNOR":
        s_att.column_dimensions[col].outlineLevel = 1
        s_att.column_dimensions[col].hidden = True

    # ------------------------------------------------------------ Power
    pwr_hdr = ["power_id", "cable_id", "freq_mhz", "value", "unit", "ambient_temp_c", "source_id", "source_ref",
               "power_w", "ln_f", "ln_p", "order_ok", "note"]
    pwr_calc = ["power_w", "ln_f", "ln_p", "order_ok"]
    write_header(s_pwr, pwr_hdr, pwr_calc)

    def pwr_formulas(r):
        return {
            "power_w": f'=IF(E{r}="kW",D{r}*1000,D{r})',
            "ln_f": f"=LN(C{r})",
            "ln_p": f"=LN(I{r})",
            "order_ok": f"=IF(B{r}=B{r - 1},C{r}>C{r - 1},COUNTIF(B$1:B{r - 1},B{r})=0)",
        }

    prow = []
    for cid, sid, unit, temp, ref, pts in PWR:
        for f, v in pts:
            prow.append((cid, f, v, unit, temp, sid, ref))
    prow.sort(key=lambda x: (order[x[0]], x[1]))
    for i, (cid, f, v, unit, temp, sid, ref) in enumerate(prow, 2):
        for j, val in enumerate([f"P{i - 1:04d}", cid, f, v, unit, temp, sid, ref], 1):
            s_pwr.cell(i, j, val).font = f_input
        fm = pwr_formulas(i)
        for j, key in enumerate(pwr_hdr, 1):
            if key in fm:
                s_pwr.cell(i, j, fm[key]).font = f_calc
        s_pwr[f"I{i}"].number_format = "#,##0"
        s_pwr[f"J{i}"].number_format = "0.000"
        s_pwr[f"K{i}"].number_format = "0.000"
    n_pwr = len(prow)
    add_table(s_pwr, "tblPower", len(pwr_hdr), n_pwr, {k: v[1:] for k, v in pwr_formulas(2).items()})
    s_pwr.conditional_formatting.add(f"L2:L{n_pwr + 1}", FormulaRule(formula=["L2=FALSE"], fill=PatternFill("solid", fgColor="FF7C80")))
    dv2 = DataValidation(type="list", formula1='"W,kW"', allow_blank=False)
    s_pwr.add_data_validation(dv2)
    dv2.add(f"E2:E{n_pwr + 500}")
    set_widths(s_pwr, {"A": 9, "B": 17, "C": 10, "D": 9, "E": 7, "F": 14, "G": 10, "H": 40, "I": 10, "J": 8, "K": 8, "L": 9, "M": 30})
    for col in "JK":
        s_pwr.column_dimensions[col].outlineLevel = 1
        s_pwr.column_dimensions[col].hidden = True

    # ------------------------------------------------------------ Cables
    calc_keys = [c[0] for c in CABLE_CALC_COLS]
    write_header(s_cab, cab_keys, calc_keys)

    def cab_formulas(r):
        A = f"A{r}"
        r0, n = f"{CC['att_first_row']}{r}", f"{CC['att_n']}{r}"
        rng = lambda col, a, b: f"INDEX(Attenuation!${col}:${col},{a}):INDEX(Attenuation!${col}:${col},{b})"
        linest = f"LINEST({rng('O', r0, f'{r0}+{n}-1')},{rng('L', r0, f'{r0}+{n}-1').replace(':INDEX(Attenuation!$L:$L', ':INDEX(Attenuation!$N:$N')},FALSE,FALSE)"
        return {
            "att_first_row": f"=MATCH({A},Attenuation!$B:$B,0)",
            "att_n": f"=COUNTIF(Attenuation!$B:$B,{A})",
            "f_min_data": f"=INDEX(Attenuation!$C:$C,{r0})",
            "f_max_data": f"=INDEX(Attenuation!$C:$C,{r0}+{n}-1)",
            "k0_fit": f"=INDEX({linest},1,3)",
            "k1_fit": f"=INDEX({linest},1,2)",
            "k2_fit": f"=INDEX({linest},1,1)",
            "fit_rms_pct": f"=SQRT(SUMIFS(Attenuation!$R:$R,Attenuation!$B:$B,{A})/{n})",
            "pwr_first_row": f'=IFERROR(MATCH({A},Power!$B:$B,0),"")',
            "pwr_n": f"=COUNTIF(Power!$B:$B,{A})",
        }

    for i, cab in enumerate(CABLES, 2):
        for j, key in enumerate(cab_keys, 1):
            if key in calc_keys:
                continue
            v = cab.get(key)
            c = s_cab.cell(i, j, v)
            c.font = f_input
        fm = cab_formulas(i)
        for key, fx in fm.items():
            c = s_cab[f"{CC[key]}{i}"]
            c.value = fx
            c.font = f_calc
        for key, fmt in (("k0_fit", "0.0000"), ("k1_fit", "0.00000"), ("k2_fit", "0.000000"), ("fit_rms_pct", "0.0%"),
                         ("k1_pub_db_100ft", "0.00000"), ("k2_pub_db_100ft", "0.00000"), ("velocity_factor", "0.00")):
            s_cab[f"{CC[key]}{i}"].number_format = fmt
    n_cab = len(CABLES)
    add_table(s_cab, "tblCables", len(cab_keys), n_cab, {k: v[1:] for k, v in cab_formulas(2).items()})
    rms = CC["fit_rms_pct"]
    s_cab.conditional_formatting.add(f"{rms}2:{rms}{n_cab + 1}",
                                     FormulaRule(formula=[f"{rms}2>0.05"], fill=PatternFill("solid", fgColor="FFE699")))
    for k_, w in {"A": 17, "B": 32, "C": 22, "D": 13, "E": 26}.items():
        s_cab.column_dimensions[k_].width = w
    for j in range(6, len(cab_keys) + 1):
        s_cab.column_dimensions[CL(j)].width = 14
    for key in ("inner_material", "dielectric_material", "shield", "jacket", "outdoor_uv", "fire_class", "connectors", "notes",
                "voltage_note", "inner_construction"):
        s_cab.column_dimensions[CC[key]].width = 28
    s_cab.column_dimensions[CC["notes"]].width = 60
    for r in range(2, n_cab + 2):
        s_cab.row_dimensions[r].height = 30
        for j in range(1, len(cab_keys) + 1):
            s_cab.cell(r, j).alignment = Alignment(vertical="top", wrap_text=True)
    s_cab.row_dimensions[1].height = 45

    # ------------------------------------------------------------ Ham_Bands
    ws = s_bands
    NROWS = 100
    ws["A1"] = "Koaksiyel kablo zayıflaması - amatör bantları"
    ws["A1"].font = f_title
    ws["A2"] = ("Değerler datasheet noktaları arasında log-log interpolasyonla hesaplanır. İtalik gri hücreler datasheet "
                "frekans aralığı dışındadır (ekstrapolasyon). En yüksek veri frekansının 2 katından yukarısı gösterilmez. "
                "Renk ölçeği her bantta kablolar arası karşılaştırmadır (yeşil = düşük kayıp).")
    ws["A2"].font = f_sub
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:AH2")
    ws.row_dimensions[2].height = 30
    fixed = [("A", "cable_id", 18), ("B", "Kablo", 32), ("C", "Z0 (Ω)", 7), ("D", "Çap (mm)", 8), ("E", "VF", 6),
             ("F", "Veri min (MHz)", 9), ("G", "Veri maks (MHz)", 9), ("H", "r0", 5), ("I", "n", 4)]
    for col, lab, w in fixed:
        ws[f"{col}5"] = lab
        ws.column_dimensions[col].width = w
        ws[f"{col}5"].font = Font(name=F, size=10, bold=True, color="FFFFFF")
        ws[f"{col}5"].fill = fill_band
        ws[f"{col}5"].alignment = center
        ws.merge_cells(f"{col}4:{col}5") if False else None
    ws.column_dimensions["H"].hidden = True
    ws.column_dimensions["I"].hidden = True
    m_cols = [CL(10 + i) for i in range(len(BANDS))]          # J..U dB/100m
    ft_cols = [CL(10 + len(BANDS) + 1 + i) for i in range(len(BANDS))]  # W..AH dB/100ft
    ws[f"{m_cols[0]}3"] = "dB / 100 m"
    ws.merge_cells(f"{m_cols[0]}3:{m_cols[-1]}3")
    ws[f"{ft_cols[0]}3"] = "dB / 100 ft (30.48 m)"
    ws.merge_cells(f"{ft_cols[0]}3:{ft_cols[-1]}3")
    for c in (f"{m_cols[0]}3", f"{ft_cols[0]}3"):
        ws[c].font = Font(name=F, size=11, bold=True, color="FFFFFF")
        ws[c].fill = fill_band
        ws[c].alignment = center
    for cols in (m_cols, ft_cols):
        for (lab, f), col in zip(BANDS, cols):
            ws[f"{col}4"] = lab
            ws[f"{col}5"] = f
            ws[f"{col}5"].number_format = '0.0" MHz"' if f < 10 else '0" MHz"'
            for rr in (4, 5):
                ws[f"{col}{rr}"].font = Font(name=F, size=10, bold=True, color="FFFFFF")
                ws[f"{col}{rr}"].fill = fill_band
                ws[f"{col}{rr}"].alignment = center
            ws.column_dimensions[col].width = 10
    ws.column_dimensions[CL(10 + len(BANDS))].width = 2
    for col in "ABCDEFGHI":
        ws[f"{col}4"].fill = fill_band
    for r in range(6, 6 + NROWS):
        k = f"ROWS($A$6:A{r})"
        ws[f"A{r}"] = f'=IF({k}>COUNTA(Cables!$A:$A)-1,"",INDEX(Cables!$A:$A,{k}+1))'
        look = lambda key: (f'=IF($A{r}="","",IF(INDEX(Cables!${CC[key]}:${CC[key]},MATCH($A{r},Cables!$A:$A,0))="","",'
                            f'INDEX(Cables!${CC[key]}:${CC[key]},MATCH($A{r},Cables!$A:$A,0))))')
        ws[f"B{r}"] = look("display_name")
        ws[f"C{r}"] = look("impedance_ohm")
        ws[f"D{r}"] = look("od_mm")
        ws[f"E{r}"] = look("velocity_factor")
        ws[f"F{r}"] = look("f_min_data")
        ws[f"G{r}"] = look("f_max_data")
        ws[f"H{r}"] = look("att_first_row")
        ws[f"I{r}"] = look("att_n")
        for col in m_cols:
            ws[f"{col}{r}"] = (f'=IF($A{r}="","",IF({col}$5>2*$G{r},"",'
                               f'{interp_formula(f"{col}$5", f"$H{r}", f"$I{r}", "Attenuation", "J", "K")}))')
            ws[f"{col}{r}"].number_format = "0.00"
        for cm, cf in zip(m_cols, ft_cols):
            ws[f"{cf}{r}"] = f'=IF({cm}{r}="","",{cm}{r}*{FT_PER_M})'
            ws[f"{cf}{r}"].number_format = "0.00"
        ws[f"E{r}"].number_format = "0.00"
        for col in [c for c, _, _ in fixed] + m_cols + ft_cols:
            ws[f"{col}{r}"].font = f_calc
            ws[f"{col}{r}"].border = box
    last = 5 + NROWS
    for cols, ref in ((m_cols, "J"), (ft_cols, ft_cols[0])):
        rng = f"{cols[0]}6:{cols[-1]}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND($A6<>"",{cols[0]}6<>"",OR({cols[0]}$5<$F6,{cols[0]}$5>$G6))'],
            font=Font(italic=True, color="404040"), stopIfTrue=False))
        for col in cols:
            ws.conditional_formatting.add(f"{col}6:{col}{last}", ColorScaleRule(
                start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="F8696B"))
    ws.freeze_panes = "C6"

    # ------------------------------------------------------------ Graph (seçilebilir 8 kablo)
    gs = s_graph
    gs["A1"] = "Zayıflama - frekans (log-log, dB/100 m)"
    gs["A1"].font = f_title
    gs["A2"] = ("Sarı hücrelerden en fazla 8 kablo seçin. Eğri, datasheet aralığının yarısından iki katına kadar çizilir "
                "(1-10000 MHz). Hesap Ham_Bands ile aynı log-log interpolasyondur.")
    gs["A2"].font = f_sub
    gs["B4"] = "Grafikteki kablolar"
    gs["B4"].font = f_bold
    defaults = ["BEL-8259", "BEL-9258", "BEL-8267", "TMS-LMR400", "MP-HYPERFLEX10", "MP-AIRBORNE10", "CS-LDF4-50A", "CS-LDF5-50A"]
    name_of = {c["cable_id"]: c["display_name"] for c in CABLES}
    NSEL = 8
    GF0, GF1 = 10, 50          # yardımcı frekans satırları (41 nokta, 1-10000 MHz)
    FX = "AA"                  # frekans sütunu
    dvg = DataValidation(type="list", formula1="Cables!$B$2:$B$500", allow_blank=True)
    gs.add_data_validation(dvg)
    for j in range(NSEL):
        cell = gs.cell(5 + j, 2, name_of.get(defaults[j]) if j < len(defaults) else None)
        cell.font = f_input
        cell.fill = fill_input
        cell.border = box
        dvg.add(cell.coordinate)
    gs.column_dimensions["A"].width = 2
    gs.column_dimensions["B"].width = 36
    gs[f"{FX}{GF0 - 1}"] = "f (MHz)"
    for k, r in enumerate(range(GF0, GF1 + 1)):
        gs[f"{FX}{r}"] = f"=10^({k}/10)"
    for j in range(NSEL):
        col = CL(28 + j)        # AB..AI
        sel = f"$B${5 + j}"

        def look(key, sel=sel):
            return f'IF({sel}="","",INDEX(Cables!${CC[key]}:${CC[key]},MATCH({sel},Cables!$B:$B,0)))'
        gs[f"{col}3"] = "=" + look("att_first_row")
        gs[f"{col}4"] = "=" + look("att_n")
        gs[f"{col}5"] = "=" + look("f_min_data")
        gs[f"{col}6"] = "=" + look("f_max_data")
        gs[f"{col}{GF0 - 1}"] = f"={sel}"
        for r in range(GF0, GF1 + 1):
            gs[f"{col}{r}"] = (f'=IF({sel}="",NA(),IF(OR(${FX}{r}<{col}$5/2,${FX}{r}>2*{col}$6),NA(),'
                               f'{interp_formula(f"${FX}{r}", f"{col}$3", f"{col}$4", "Attenuation", "J", "K")}))')
    for j in range(NSEL + 1):
        gs.column_dimensions[CL(27 + j)].hidden = True
    ch = ScatterChart()
    ch.title = None
    ch.style = 2
    ch.height = 15
    ch.width = 24
    ch.x_axis.title = "Frekans (MHz)"
    ch.y_axis.title = "dB / 100 m"
    ch.x_axis.scaling.logBase = 10
    ch.y_axis.scaling.logBase = 10
    ch.x_axis.scaling.min = 1
    ch.x_axis.scaling.max = 10000
    ch.y_axis.scaling.min = 0.1
    ch.y_axis.scaling.max = 300
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    ch.legend.overlay = False
    ch.x_axis.crosses = "min"
    ch.y_axis.crosses = "min"
    ch.x_axis.number_format = "0"
    ch.y_axis.number_format = "0.0#"
    ch.visible_cells_only = False
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.data_source import StrRef
    xref = Reference(gs, min_col=27, min_row=GF0, max_row=GF1)
    for j in range(NSEL):
        yref = Reference(gs, min_col=28 + j, min_row=GF0, max_row=GF1)
        se = Series(yref, xref, title_from_data=False)
        se.tx = SeriesLabel(strRef=StrRef(f"'Graph'!$B${5 + j}"))
        se.marker.symbol = "none"
        se.smooth = False
        ch.series.append(se)
    gs.add_chart(ch, "D4")

    # ------------------------------------------------------------ Calculator
    cs = s_calc
    cs["B2"] = "Kablo kaybı hesaplayıcı"
    cs["B2"].font = f_title
    cs["B3"] = "Sarı hücreleri değiştirin. Diğer hücreler formüldür."
    cs["B3"].font = f_sub
    default_cable = next(c["display_name"] for c in CABLES if c["cable_id"] == "TMS-LMR400")
    inputs = [("B5", "Kablo", "C5", default_cable), ("B6", "Frekans (MHz)", "C6", 145),
              ("B7", "Kablo uzunluğu (m)", "C7", 20), ("B8", "Verici çıkış gücü (W)", "C8", 100),
              ("B9", "Anten (yük) SWR", "C9", 1.5)]
    for lb, text, vc, val in inputs:
        cs[lb] = text
        cs[lb].font = f_bold
        cs[vc] = val
        cs[vc].font = f_input
        cs[vc].fill = fill_input
        cs[vc].border = box
    dvc = DataValidation(type="list", formula1="Cables!$B$2:$B$500", allow_blank=False)
    cs.add_data_validation(dvc)
    dvc.add("C5")
    dvn = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    cs.add_data_validation(dvn)
    dvn.add("C6:C8")
    dvs = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="1")
    cs.add_data_validation(dvs)
    dvs.add("C9")
    # yardımcılar (F-G)
    cs["F4"] = "Yardımcı değerler"
    cs["F4"].font = f_bold
    helpers = [("cable_id", f"=INDEX(Cables!$A:$A,MATCH(C5,Cables!$B:$B,0))"),
               ("att_first_row", f"=INDEX(Cables!${CC['att_first_row']}:${CC['att_first_row']},MATCH(G5,Cables!$A:$A,0))"),
               ("att_n", f"=INDEX(Cables!${CC['att_n']}:${CC['att_n']},MATCH(G5,Cables!$A:$A,0))"),
               ("pwr_first_row", f"=INDEX(Cables!${CC['pwr_first_row']}:${CC['pwr_first_row']},MATCH(G5,Cables!$A:$A,0))"),
               ("pwr_n", f"=INDEX(Cables!${CC['pwr_n']}:${CC['pwr_n']},MATCH(G5,Cables!$A:$A,0))"),
               ("velocity_factor", f"=INDEX(Cables!${CC['velocity_factor']}:${CC['velocity_factor']},MATCH(G5,Cables!$A:$A,0))"),
               ("f_min_data", f"=INDEX(Cables!${CC['f_min_data']}:${CC['f_min_data']},MATCH(G5,Cables!$A:$A,0))"),
               ("f_max_data", f"=INDEX(Cables!${CC['f_max_data']}:${CC['f_max_data']},MATCH(G5,Cables!$A:$A,0))"),
               ("Γ yük", "=(C9-1)/(C9+1)"),
               ("a = 10^(ML/10)", "=10^(C15/10)")]
    for i, (lab, fx) in enumerate(helpers, 5):
        cs[f"F{i}"] = lab
        cs[f"F{i}"].font = Font(name=F, size=9, color="808080")
        cs[f"G{i}"] = fx
        cs[f"G{i}"].font = Font(name=F, size=9, color="808080")
    outs = [
        (13, "Zayıflama (dB/100 m)", "=" + interp_formula("C6", "G6", "G7", "Attenuation", "J", "K"), "0.000"),
        (14, "Zayıflama (dB/100 ft)", f"=C13*{FT_PER_M}", "0.000"),
        (15, "Hat kaybı - uyumlu yük (dB)", "=C13*C7/100", "0.00"),
        (16, "SWR kaynaklı ek kayıp (dB)", "=C17-C15", "0.00"),
        (17, "Toplam hat kaybı (dB)", "=10*LOG10((G14^2-G13^2)/(G14*(1-G13^2)))", "0.00"),
        (18, "Antene ulaşan güç (W)", "=C8*10^(-C17/10)", "0.0"),
        (19, "Hat verimi", "=C18/C8", "0.0%"),
        (20, "Verici ucundaki SWR", "=(1+G13/G14)/(1-G13/G14)", "0.00"),
        (21, "Maks. ortalama güç @ f (W)", '=IF(G9<2,"veri yok",' + interp_formula("C6", "G8", "G9", "Power", "J", "K") + ")", "#,##0"),
        (22, "Güç durumu", '=IF(ISNUMBER(C21),IF(C8>C21,"UYARI: giriş gücü sınırın üzerinde","Uygun"),"Datasheet güç vermiyor")', "@"),
        (23, "Kablo içi dalga boyu λ (m)", '=IF(N(G10)=0,"VF yayımlanmamış",299.792458/C6*G10)', "0.000"),
        (24, "λ/4 fiziksel uzunluk (m)", '=IF(ISNUMBER(C23),C23/4,"-")', "0.000"),
        (25, "Hattın elektriksel uzunluğu (λ)", '=IF(ISNUMBER(C23),C7/C23,"-")', "0.000"),
        (26, "Veri durumu", '=IF(OR(C6<G11,C6>G12),"Ekstrapolasyon: datasheet aralığı "&G11&"-"&G12&" MHz dışında","Datasheet aralığında (interpolasyon)")', "@"),
    ]
    for r, lab, fx, fmt in outs:
        cs[f"B{r}"] = lab
        cs[f"B{r}"].font = f_bold
        cs[f"C{r}"] = fx
        cs[f"C{r}"].font = f_calc
        cs[f"C{r}"].fill = fill_out
        cs[f"C{r}"].border = box
        cs[f"C{r}"].number_format = fmt
    cs["B12"] = "Sonuçlar"
    cs["B12"].font = Font(name=F, size=11, bold=True, color="1F3864")
    cs["B27"] = ("Toplam kayıp: ARRL Antenna Book uyumsuz hat formülü, TL = 10·log10[(a² − |Γ|²) / (a·(1 − |Γ|²))], a = 10^(ML/10). "
                 "Güç sınırı datasheet ortam sıcaklığı ve uyumlu yük içindir; SWR ve yüksek sıcaklık sınırı düşürür.")
    cs["B27"].font = f_sub
    cs["B27"].alignment = wrap
    cs.merge_cells("B27:I27")
    cs.row_dimensions[27].height = 42
    # karşılaştırma tablosu
    cs["B29"] = "Tüm kablolar - aynı frekans, uzunluk, güç ve SWR ile"
    cs["B29"].font = Font(name=F, size=11, bold=True, color="1F3864")
    cmp_hdr = ["Kablo", "dB/100 m", "Toplam kayıp (dB)", "Antene güç (W)", "Verim", "Maks. güç (W)", "Not"]
    for j, h in enumerate(cmp_hdr):
        c = cs.cell(30, 2 + j, h)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = fill_band
        c.alignment = center
    for r in range(31, 31 + NROWS):
        k = f"ROWS($B$31:B{r})+1"
        ok = f'ROWS($B$31:B{r})>COUNTA(Cables!$A:$A)-1'
        g = lambda key: f"INDEX(Cables!${CC[key]}:${CC[key]},{k})"
        cs[f"B{r}"] = f'=IF({ok},"",{g("display_name")})'
        cs[f"C{r}"] = f'=IF(B{r}="","",{interp_formula("$C$6", g("att_first_row"), g("att_n"), "Attenuation", "J", "K")})'
        ml = f"(C{r}*$C$7/100)"
        a = f"10^({ml}/10)"
        cs[f"D{r}"] = f'=IF(B{r}="","",10*LOG10(({a}^2-$G$13^2)/({a}*(1-$G$13^2))))'
        cs[f"E{r}"] = f'=IF(B{r}="","",$C$8*10^(-D{r}/10))'
        cs[f"F{r}"] = f'=IF(B{r}="","",E{r}/$C$8)'
        cs[f"G{r}"] = (f'=IF(B{r}="","",IF({g("pwr_n")}<2,"-",'
                       f'{interp_formula("$C$6", g("pwr_first_row"), g("pwr_n"), "Power", "J", "K")}))')
        cs[f"H{r}"] = (f'=IF(B{r}="","",MID(IF(OR($C$6<{g("f_min_data")},$C$6>{g("f_max_data")})," / ekstrapolasyon","")'
                       f'&IF(AND(ISNUMBER(G{r}),$C$8>G{r})," / güç sınırı aşılıyor","")'
                       f'&IF({g("impedance_ohm")}<>50," / Z0="&{g("impedance_ohm")}&" Ω",""),4,200))')
        for col, fmt in (("C", "0.00"), ("D", "0.00"), ("E", "0.0"), ("F", "0.0%"), ("G", "#,##0")):
            cs[f"{col}{r}"].number_format = fmt
        for col in "BCDEFGH":
            cs[f"{col}{r}"].font = f_calc
            cs[f"{col}{r}"].border = box
    cs.conditional_formatting.add(f"F31:F{30 + NROWS}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                                                                     mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    set_widths(cs, {"A": 2, "B": 34, "C": 26, "D": 16, "E": 14, "F": 16, "G": 14, "H": 34, "I": 4})
    cs.column_dimensions["G"].width = 14

    # ------------------------------------------------------------ Legacy
    lwb = openpyxl.load_workbook(LEGACY, data_only=True)
    l1 = lwb["Sayfa1 (2)"]
    l2 = lwb["Sayfa2"]
    lix_hdr = ["legacy_key", "legacy_sheet", "legacy_row", "legacy_name", "z0_old", "od_old_mm", "cable_id", "match_type",
               "status", "note", "z0_new", "od_new_mm"]
    write_header(s_lix, lix_hdr, ["z0_new", "od_new_mm"])
    lpt_hdr = ["legacy_key", "legacy_name", "band", "f_lo_mhz", "f_hi_mhz", "old_db_100ft", "cable_id", "new_lo_db_100ft",
               "new_hi_db_100ft", "dev_pct", "status"]
    write_header(s_lpt, lpt_hdr, ["legacy_name", "cable_id", "new_lo_db_100ft", "new_hi_db_100ft", "dev_pct", "status"])
    li = 2
    lp = 2
    entries = []
    for r in range(5, l1.max_row + 1):
        name = l1[f"B{r}"].value
        if not name:
            continue
        vals = {col: l1[f"{col}{r}"].value for col, *_ in LEGACY_BANDS}
        entries.append(("Sayfa1 (2)", r, name, l1[f"C{r}"].value, l1[f"D{r}"].value, vals, ""))
    # Sayfa2: başlık yok; değerler bir sütun sağa kaymış görünüyor (E=çap, F=1-10 MHz ...)
    for r in range(2, l2.max_row + 1):
        name = l2[f"B{r}"].value
        if not name:
            continue
        cols = [b[0] for b in LEGACY_BANDS]
        vals = {}
        for k, col in enumerate(cols):
            src_col = CL(openpyxl.utils.column_index_from_string(col) + 1)
            vals[col] = l2[f"{src_col}{r}"].value
        entries.append(("Sayfa2", r, name, l2[f"C{r}"].value, l2[f"E{r}"].value, vals,
                        "75 Ω tablo: başlık yok, değerler bir sütun kaymış varsayıldı (E=çap, F=1-10 MHz, ...)."))
    for sheet, r, name, z0, od, vals, extra in entries:
        key = f"{'S1' if sheet.startswith('Sayfa1') else 'S2'}-R{r:02d}"
        cid, mtype, note = ("", "", "")
        if sheet.startswith("Sayfa1") and r in LEGACY_MAP:
            cid, mtype, note = LEGACY_MAP[r]
        if sheet == "Sayfa2" and r in LEGACY_MAP_S2:
            cid, mtype, note = LEGACY_MAP_S2[r]
        if sheet.startswith("Sayfa1") and r in LEGACY_NOTES:
            note = (note + " " + LEGACY_NOTES[r]).strip()
        note = (note + " " + extra).strip()
        has_vals = any(v is not None for v in vals.values())
        if sheet.startswith("Sayfa1") and r in LEGACY_NO_SOURCE:
            status = "Kaynak bulunamadı"
        elif cid:
            status = "Doğrulandı" if has_vals else "Eşlendi (eski veri yok)"
        else:
            status = "Veri yok" if not has_vals else "Bekliyor"
        row = [key, sheet, r, name, z0, od, cid or None, mtype or None, status, note or None]
        for j, v in enumerate(row, 1):
            s_lix.cell(li, j, v).font = f_input
        s_lix.cell(li, 11, f'=IF(G{li}="","",INDEX(Cables!${CC["impedance_ohm"]}:${CC["impedance_ohm"]},MATCH(G{li},Cables!$A:$A,0)))').font = f_calc
        s_lix.cell(li, 12, f'=IF(G{li}="","",INDEX(Cables!${CC["od_mm"]}:${CC["od_mm"]},MATCH(G{li},Cables!$A:$A,0)))').font = f_calc
        li += 1
        for col, lab, flo, fhi in LEGACY_BANDS:
            v = vals.get(col)
            if v is None:
                continue
            s_lpt.cell(lp, 1, key).font = f_input
            s_lpt.cell(lp, 2, f"=INDEX(Legacy_Index!$D:$D,MATCH(A{lp},Legacy_Index!$A:$A,0))").font = f_calc
            for j, val in ((3, lab), (4, flo), (5, fhi), (6, v)):
                s_lpt.cell(lp, j, val).font = f_input
            s_lpt.cell(lp, 7, f'=IF(INDEX(Legacy_Index!$G:$G,MATCH(A{lp},Legacy_Index!$A:$A,0))="","",'
                              f'INDEX(Legacy_Index!$G:$G,MATCH(A{lp},Legacy_Index!$A:$A,0)))').font = f_calc
            r0 = f"INDEX(Cables!${CC['att_first_row']}:${CC['att_first_row']},MATCH(G{lp},Cables!$A:$A,0))"
            nn = f"INDEX(Cables!${CC['att_n']}:${CC['att_n']},MATCH(G{lp},Cables!$A:$A,0))"
            s_lpt.cell(lp, 8, f'=IF(G{lp}="","",{interp_formula(f"D{lp}", r0, nn, "Attenuation", "J", "K")}*{FT_PER_M})').font = f_calc
            s_lpt.cell(lp, 9, f'=IF(G{lp}="","",{interp_formula(f"E{lp}", r0, nn, "Attenuation", "J", "K")}*{FT_PER_M})').font = f_calc
            s_lpt.cell(lp, 10, f'=IF(G{lp}="","",IF(F{lp}<H{lp},F{lp}/H{lp}-1,IF(F{lp}>I{lp},F{lp}/I{lp}-1,0)))').font = f_calc
            fmn = f"INDEX(Cables!${CC['f_min_data']}:${CC['f_min_data']},MATCH(G{lp},Cables!$A:$A,0))"
            fmx = f"INDEX(Cables!${CC['f_max_data']}:${CC['f_max_data']},MATCH(G{lp},Cables!$A:$A,0))"
            s_lpt.cell(lp, 11, f'=IF(G{lp}="","Doğrulanmadı",IF(ABS(J{lp})<=0.1,"Uyumlu (±%10)",'
                               f'IF(ABS(J{lp})<=0.25,"Fark %10-25","Fark >%25"))'
                               f'&IF(OR(D{lp}<{fmn},E{lp}>{fmx})," · ekstrapolasyon",""))').font = f_calc
            for col, fmt in (("H", "0.00"), ("I", "0.00"), ("J", "+0%;-0%;0%")):
                s_lpt[f"{col}{lp}"].number_format = fmt
            lp += 1
    add_table(s_lix, "tblLegacyIndex", len(lix_hdr), li - 2)
    add_table(s_lpt, "tblLegacyPoints", len(lpt_hdr), lp - 2)
    s_lpt.conditional_formatting.add(f"K2:K{lp - 1}", FormulaRule(formula=['ISNUMBER(SEARCH("Fark >%25",K2))'], fill=PatternFill("solid", fgColor="F8CBAD")))
    s_lpt.conditional_formatting.add(f"K2:K{lp - 1}", FormulaRule(formula=['ISNUMBER(SEARCH("Fark %10-25",K2))'], fill=PatternFill("solid", fgColor="FFE699")))
    s_lpt.conditional_formatting.add(f"K2:K{lp - 1}", FormulaRule(formula=['ISNUMBER(SEARCH("Uyumlu",K2))'], fill=PatternFill("solid", fgColor="C6EFCE")))
    s_lix.conditional_formatting.add(f"I2:I{li - 1}", FormulaRule(formula=['LEFT(I2,4)="Doğr"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    s_lix.conditional_formatting.add(f"I2:I{li - 1}", FormulaRule(formula=['I2="Kaynak bulunamadı"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    set_widths(s_lix, {"A": 10, "B": 11, "C": 8, "D": 42, "E": 8, "F": 9, "G": 17, "H": 12, "I": 18, "J": 70, "K": 8, "L": 10})
    set_widths(s_lpt, {"A": 10, "B": 40, "C": 12, "D": 9, "E": 9, "F": 11, "G": 17, "H": 14, "I": 14, "J": 9, "K": 16})
    for r in range(2, li):
        s_lix[f"J{r}"].alignment = wrap

    # ------------------------------------------------------------ Data_Dictionary
    dd_hdr = ["table", "column", "description_tr", "unit", "kind"]
    write_header(s_dd, dd_hdr)
    dd = []
    for k, d, u in CABLE_COLS:
        dd.append(("Cables", k, d, u, "girdi"))
    for k, d, u in CABLE_CALC_COLS:
        dd.append(("Cables", k, d, u, "formül"))
    att_desc = {
        "point_id": ("Nokta kimliği", ""), "cable_id": ("Cables.cable_id", ""), "freq_mhz": ("Frekans", "MHz"),
        "value": ("Datasheet'te yazan değer (olduğu gibi)", "unit sütunu"), "unit": ("Değerin birimi: dB/100m veya dB/100ft", ""),
        "temp_c": ("Ölçüm/spesifikasyon sıcaklığı (belirtilmişse)", "°C"), "source_id": ("Sources.source_id", ""),
        "source_ref": ("Belgedeki yer (tablo/sayfa)", ""), "att_db_100m": ("Zayıflama, dB/100 m'ye çevrilmiş", "dB/100m"),
        "ln_f": ("ln(f) - interpolasyon yardımcı", ""), "ln_att": ("ln(α) - interpolasyon yardımcı", ""),
        "w0": ("QA regresyonu yardımcı: 1/α", ""), "w1": ("QA regresyonu yardımcı: √f/α", ""), "w2": ("QA regresyonu yardımcı: f/α", ""),
        "one": ("QA regresyonu yardımcı: 1", ""), "model_db_100m": ("QA modeli k0+k1√f+k2f değeri", "dB/100m"),
        "dev_pct": ("Model / datasheet - 1 (veri girişi kontrolü)", "%"), "dev_sq": ("dev_pct²", ""),
        "order_ok": ("Satırlar kablo bazında gruplu ve frekansa göre artan mı", ""), "note": ("Not", ""),
    }
    for k in att_hdr:
        d, u = att_desc[k]
        dd.append(("Attenuation", k, d, u, "formül" if k in att_calc else "girdi"))
    pwr_desc = {"power_id": ("Nokta kimliği", ""), "cable_id": ("Cables.cable_id", ""), "freq_mhz": ("Frekans", "MHz"),
                "value": ("Datasheet'te yazan ortalama (CW) güç", "unit sütunu"), "unit": ("W veya kW", ""),
                "ambient_temp_c": ("Güç değerinin ortam sıcaklığı", "°C"), "source_id": ("Sources.source_id", ""),
                "source_ref": ("Belgedeki yer", ""), "power_w": ("Güç, W'a çevrilmiş", "W"), "ln_f": ("ln(f)", ""),
                "ln_p": ("ln(P)", ""), "order_ok": ("Sıra kontrolü", ""), "note": ("Not", "")}
    for k in pwr_hdr:
        d, u = pwr_desc[k]
        dd.append(("Power", k, d, u, "formül" if k in pwr_calc else "girdi"))
    for k in src_hdr:
        dd.append(("Sources", k, {"source_id": "Kaynak kimliği", "publisher": "Yayımlayan üretici", "document_title": "Belge başlığı",
                                  "part_number": "Parça no", "revision": "Revizyon", "document_date": "Belge tarihi",
                                  "host": "Belgeyi barındıran (üretici / distribütör kopyası)", "url": "İndirme adresi",
                                  "accessed": "Erişim tarihi", "local_file": "Proje klasöründeki arşiv kopyası",
                                  "notes": "Not"}[k], "", "girdi"))
    for i, row in enumerate(dd, 2):
        for j, v in enumerate(row, 1):
            s_dd.cell(i, j, v).font = f_base
    add_table(s_dd, "tblDataDictionary", len(dd_hdr), len(dd))
    set_widths(s_dd, {"A": 14, "B": 24, "C": 70, "D": 12, "E": 9})

    # ------------------------------------------------------------ README
    rd = s_readme
    rd.column_dimensions["A"].width = 2
    rd.column_dimensions["B"].width = 120
    lines = [
        ("TA3HX Koaksiyel Kablo Veritabanı", f_title),
        (f"Sürüm 1.0 · 2026-10-09 · {n_cab} kablo · {n_att} zayıflama noktası · {n_pwr} güç noktası · {len(SOURCES) - 1} kaynak belge "
         f"(PDF/metin kopyaları 'kaynaklar' klasöründe)", f_sub),
        ("", None),
        ("AMAÇ", f_bold),
        ("Her değeri bir datasheet'e bağlanmış, hesapları formülle yapılan ve ileride bir programa veritabanı olarak aktarılabilecek "
         "koaksiyel kablo tablosu. Orijinal elle hazırlanmış tablo Legacy_* sayfalarında korunur ve yeni verilerle karşılaştırılır.", None),
        ("", None),
        ("SAYFALAR", f_bold),
        ("Ham_Bands - Rapor: her kablo için amatör bantlarında dB/100 m ve dB/100 ft. Cables sayfasına eklenen kablo otomatik görünür (100 satıra kadar).", None),
        ("Graph - Seçilen en fazla 8 kablonun 1-10000 MHz log-log zayıflama eğrisi (sarı hücrelerden seçim).", None),
        ("Calculator - Kablo, frekans, uzunluk, güç ve SWR girin: kayıp, antene ulaşan güç, güç sınırı, dalga boyu. Altta tüm kabloların karşılaştırması.", None),
        ("Cables - Kablo ana tablosu (bir satır = bir ürün). Elektriksel, mekanik, güç ve uygunluk parametreleri. Gri başlıklı sütunlar formüldür.", None),
        ("Attenuation - Zayıflama noktaları (bir satır = bir datasheet değeri). Değer datasheet'te yazdığı birimde girilir, dB/100 m'ye formülle çevrilir.", None),
        ("Power - Ortalama (CW) güç taşıma noktaları, datasheet'te yazdığı birimde.", None),
        ("Sources - Kaynak belgeler: URL, revizyon, erişim tarihi ve 'kaynaklar' klasöründeki arşiv PDF'i.", None),
        ("Legacy_Index / Legacy_Points - Eski tablo satırları, yeni kablo kimlikleriyle eşleme ve bant bazında karşılaştırma.", None),
        ("Data_Dictionary - Tüm sütunların Türkçe açıklaması ve birimi (program geliştirirken referans).", None),
        ("", None),
        ("HESAPLAMA YÖNTEMİ", f_bold),
        ("1) Ana hesap - log-log interpolasyon: Hedef frekansın iki yanındaki datasheet noktaları arasında ln(α)-ln(f) düzleminde doğrusal "
         "interpolasyon yapılır. Datasheet frekanslarında sonuç datasheet ile birebir aynıdır. Eski tablodaki bant aralığı (ör. 148-174 MHz) "
         "belirsizliği ortadan kalkar; her değer tek bir frekansa aittir.", None),
        ("2) Ekstrapolasyon: Veri aralığı dışında uçtaki segmentin eğimi (α ∝ f^n) kullanılır. Ham_Bands'te italik gri gösterilir; en yüksek "
         "veri frekansının 2 katından yukarısı gösterilmez.", None),
        ("3) Kalite kontrolü (QA) - fiziksel model: α(f) = k0 + k1·√f + k2·f. k1·√f iletken (deri etkisi) kaybı, k2·f dielektrik kaybıdır. "
         "Katsayılar her kablo için göreli hatayı en aza indiren ağırlıklı en küçük kareler (LINEST) ile Cables sayfasında hesaplanır. "
         "Attenuation.dev_pct sütunu her noktanın modelden sapmasını gösterir: >%5 sarı, >%10 kırmızı -> veri girişini kontrol edin.", None),
        ("   Not: Times Microwave yalnızca 2 terimli modeli (k0=0) yayımlar (Cables.k1_pub / k2_pub, dB/100 ft). HF'de ince ve köpük "
         "dielektrikli kablolarda 2 terimli model %20-30 sapabildiği için ana hesapta interpolasyon kullanılır.", None),
        ("4) Güç: Power noktaları arasında aynı log-log interpolasyon. Datasheet ortam sıcaklığı (genelde 40 °C) ve uyumlu yük içindir.", None),
        ("5) SWR'li toplam kayıp (Calculator): ARRL Antenna Book formülü TL = 10·log10[(a² − |Γ|²)/(a·(1 − |Γ|²))], a = 10^(ML/10).", None),
        ("Birimler: dB/100 ft = dB/100 m × 0.3048. Frekans MHz. Sıcaklık: datasheet koşulları (Times 25 °C, M&P ve CommScope 20 °C, Belden belirtilmemiş).", None),
        ("", None),
        ("RENK KODLARI", f_bold),
        ("Mavi yazı = elle girilen datasheet değeri · Siyah yazı = formül · Gri başlık = formül sütunu (değiştirmeyin) · Sarı dolgu = Calculator girdileri.", None),
        ("", None),
        ("YENİ KABLO EKLEME", f_bold),
        ("1) Sources: belgeyi yeni bir source_id ile ekleyin, PDF'i 'kaynaklar' klasörüne kaydedin.", None),
        ("2) Cables: tablonun altına yeni satır ekleyin (cable_id benzersiz olmalı). Formül sütunları tabloyla otomatik uzar.", None),
        ("3) Attenuation ve Power: noktaları ekleyin. ÖNEMLİ: satırlar kablo bazında bitişik ve frekansa göre artan sırada olmalı "
         "(interpolasyon buna dayanır). Eklemeden sonra Veri > Sırala: cable_id, sonra freq_mhz. order_ok sütunu YANLIŞ ise sıra bozuktur.", None),
        ("4) Legacy_Index: eski tablodaki karşılık gelen satıra cable_id ve eşleşme türünü yazın; Legacy_Points karşılaştırması kendiliğinden güncellenir.", None),
        ("5) En az 2 nokta interpolasyon, en az 4 nokta QA modeli için gerekir.", None),
        ("", None),
        ("VERİTABANINA AKTARMA", f_bold),
        ("Cables, Attenuation, Power, Sources tabloları ilişkisel yapıdadır (cable_id ve source_id anahtarları). Hazır dışa aktarım: "
         "export/coax.sqlite (yabancı anahtarlı SQLite) ve export/csv/*.csv (UTF-8, virgül ayraçlı, ondalık nokta). ham_bands tablosu "
         "Ham_Bands sayfasının uzun biçimidir. Excel'de değişiklik yaptıktan sonra kaydedip 'python -X utf8 scripts/export_db.py' ile yenileyin.", None),
        ("", None),
        ("SINIRLAMALAR", f_bold),
        ("Değer türü üreticiye göre değişir: çoğu nominal/tipik; Reçber (RWC, RG 213 U) ve Ericsson TZC 500 32 MAKSİMUM değer verir "
         "(Cables.notes). Belden H155/H1000 maksimumu nominalin %10 üstü; CommScope ±%5 garanti verir. Konnektör ve eklem kayıpları dahil değildir.", None),
        ("Kaynak niteliği: Sources.host sütunu belgenin üreticiden mi, distribütör/satıcı kopyasından mı alındığını gösterir. Andrew/CommScope "
         "HELIAX verileri 2002-2007 Andrew datasheet'lerinden; bu ürünlerin bir kısmı üretimden kalkmış olabilir.", None),
        ("Eski tablo: 60 satırın 59'u bir kabloyla eşlendi. 'RG58 CellFoil (9006)' için yayımlanmış veri bulunamadı. Aynı kablonun tekrarları "
         "(ör. 1/2\" LDF ve 1/2\" HELIAX) aynı cable_id'ye bağlandı. Bant aralığı dışında kalan eski değerler Legacy_Points'te renkli işaretlidir.", None),
    ]
    for i, (txt, font) in enumerate(lines, 1):
        c = rd.cell(i, 2, txt)
        c.font = font or f_base
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # genel font (stil verilmemiş hücreler)
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for c in row:
                if c.font is None or c.font.name != F:
                    c.font = Font(name=F, size=c.font.size if c.font else 10, bold=c.font.b, italic=c.font.i,
                                  color=c.font.color)
    wb.save(OUT)
    print("yazıldı:", OUT, "| kablo", n_cab, "| zayıflama", n_att, "| güç", n_pwr, "| legacy satır", li - 2, "| legacy nokta", lp - 2)


if __name__ == "__main__":
    build()
